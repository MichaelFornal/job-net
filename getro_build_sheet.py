#!/usr/bin/env python3
"""
getro_build_sheet.py
Aggregate getro_all_jobs.jsonl (raw per-board rows from getro_fetch_jobs.py /
getro_fetch_pdl.py) into a deduplicated master CSV + organized Excel workbook.

Dedupe, two passes:
  1. exact Getro job id (same record served to multiple networks)
  2. same (company, title, apply_url) — or (company, title, locations) when no
     URL — across boards; board lists merged, earliest posting date kept.

Outputs (this directory):
  getro_jobs_master.csv          full deduped table (always complete)
  Getro_All_Jobs_<date>.xlsx     README / All Jobs / Boards / Companies
  getro_jobs_pdl_only.csv        only if the deduped total exceeds Excel's row
                                 limit: jobs listed solely on People Data Labs
                                 move here and the workbook keeps the rest.

Memory notes: rows are fixed-index lists (not dicts) and every repeated string
is interned via one memo dict — ~2.15M raw rows fit a 16GB machine.
"""

import csv
import json
import os
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


def safe(v):
    """Strip XML-illegal control chars (openpyxl raises on them)."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

BASE = os.path.dirname(os.path.abspath(__file__))
IN_JSONL = os.path.join(BASE, "getro_all_jobs.jsonl")
BOARDS_CSV = os.path.join(BASE, "getro_job_boards.csv")
STATE_JSON = os.path.join(BASE, "getro_fetch_state.json")
OUT_CSV = os.path.join(BASE, "getro_jobs_master.csv")
PDL_CSV = os.path.join(BASE, "getro_jobs_pdl_only.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Getro_All_Jobs_{SNAPSHOT}.xlsx")

XLSX_ROW_CAP = 1_048_000
PDL_BOARD = "People Data Labs"

COLUMNS = [
    "title", "company", "location", "work_mode", "seniority", "posted",
    "salary_min", "salary_max", "salary_currency", "salary_period", "equity",
    "skills", "company_stage", "company_size", "company_industry",
    "apply_url", "board_job_url", "boards", "board_count",
    "source", "getro_job_id",
]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)
ORG_SLUG, BOARDS, TS = N, N + 1, N + 2  # aux tail slots

_MEMO = {}


def I(s):  # intern
    return _MEMO.setdefault(s, s)


def pretty(s):
    return I(str(s).replace("_", " ").strip().title()) if s else ""


WM = {"on_site": "On-site", "hybrid": "Hybrid", "remote": "Remote", None: "",
      "": ""}


def norm(rec):
    """One raw JSONL line -> fixed-index row list."""
    j = rec["job"]
    org = j.get("organization") or {}
    locs = j.get("locations") or []
    if not locs and j.get("searchable_locations"):
        locs = j["searchable_locations"][:1]
    ts = j.get("created_at")
    ts = ts if isinstance(ts, (int, float)) and ts > 0 else 0
    posted = I(datetime.fromtimestamp(ts, tz=timezone.utc)
               .strftime("%Y-%m-%d")) if ts else ""
    url = (j.get("url") or "").strip()
    slug = j.get("slug") or ""
    public = j.get("compensation_public")
    mn = j.get("compensation_amount_min_cents")
    mx = j.get("compensation_amount_max_cents")
    equity = j.get("compensation_offers_equity")
    row = [""] * (N + 3)
    row[IDX["title"]] = safe((j.get("title") or "").strip())
    row[IDX["company"]] = I((org.get("name") or "").strip())
    row[IDX["location"]] = I("; ".join(x for x in locs if x))
    row[IDX["work_mode"]] = I(WM.get(j.get("work_mode"),
                                     pretty(j.get("work_mode"))))
    row[IDX["seniority"]] = pretty(j.get("seniority"))
    row[IDX["posted"]] = posted
    row[IDX["salary_min"]] = round(mn / 100) if public and isinstance(
        mn, (int, float)) else ""
    row[IDX["salary_max"]] = round(mx / 100) if public and isinstance(
        mx, (int, float)) else ""
    row[IDX["salary_currency"]] = I(j.get("compensation_currency") or "") \
        if public else ""
    row[IDX["salary_period"]] = pretty(j.get("compensation_period")).replace(
        "Period Not Defined", "") if public else ""
    row[IDX["equity"]] = {True: "Yes", False: "No"}.get(equity, "")
    row[IDX["skills"]] = I(", ".join((j.get("skills") or [])[:8]))
    row[IDX["company_stage"]] = pretty(org.get("stage"))
    row[IDX["company_size"]] = org.get("head_count") or ""
    row[IDX["company_industry"]] = I(", ".join(
        (org.get("industry_tags") or [])[:4]))
    row[IDX["apply_url"]] = url
    row[IDX["board_job_url"]] = (rec["board_url"].rstrip("/") + "/jobs/"
                                 + slug) if slug else ""
    row[IDX["source"]] = pretty(j.get("source"))
    row[IDX["getro_job_id"]] = j.get("id")
    row[ORG_SLUG] = I(org.get("slug") or "")
    row[BOARDS] = {I(rec["board_name"])}
    row[TS] = ts
    return row


MERGE_IF_EMPTY = [IDX[k] for k in ("apply_url", "location", "skills",
                                   "company_industry", "salary_min",
                                   "salary_max", "seniority")]


def merge(a, b):
    a[BOARDS] |= b[BOARDS]
    if b[TS] and (not a[TS] or b[TS] < a[TS]):
        a[TS], a[IDX["posted"]] = b[TS], b[IDX["posted"]]
    for i in MERGE_IF_EMPTY:
        if a[i] == "" and b[i] != "":
            a[i] = b[i]


def main():
    # ---- pass 1: stream + dedupe by getro job id
    by_id = {}
    raw_rows = 0
    bad_lines = 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            raw_rows += 1
            if raw_rows % 250_000 == 0:
                print(f"  ...{raw_rows} raw rows", flush=True)
            try:
                r = norm(json.loads(line))
            except (json.JSONDecodeError, KeyError):
                bad_lines += 1  # truncated tail from a killed writer
                continue
            jid = r[IDX["getro_job_id"]]
            prev = by_id.get(jid)
            if prev is not None:
                merge(prev, r)
            else:
                by_id[jid] = r
    after_id = len(by_id)

    # ---- pass 2: same company+title+url (or +location when no url)
    by_key = {}
    for r in by_id.values():
        key = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               r[IDX["apply_url"]] or "loc:" + r[IDX["location"]].lower())
        prev = by_key.get(key)
        if prev is not None:
            merge(prev, r)
        else:
            by_key[key] = r
    by_id = None
    rows = sorted(by_key.values(), key=lambda r: r[TS], reverse=True)
    by_key = None
    for r in rows:
        names = sorted(r[BOARDS])
        r[IDX["board_count"]] = len(names)
        r[IDX["boards"]] = "; ".join(names[:6]) + (
            f" (+{len(names) - 6} more)" if len(names) > 6 else "")
    print(f"raw rows {raw_rows} -> unique ids {after_id} "
          f"-> deduped {len(rows)}"
          + (f"  ({bad_lines} malformed lines skipped)" if bad_lines else ""),
          flush=True)

    # ---- master CSV (always complete)
    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print(f"wrote {OUT_CSV}", flush=True)

    # ---- split decision for the workbook
    pdl_only = []
    eco = rows
    split = len(rows) > XLSX_ROW_CAP
    if split:
        eco, pdl_only = [], []
        for r in rows:
            (pdl_only if r[BOARDS] == {PDL_BOARD} else eco).append(r)
        with open(PDL_CSV, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(COLUMNS)
            for r in pdl_only:
                w.writerow(r[:N])
        print(f"split: {len(eco)} ecosystem rows in workbook, "
              f"{len(pdl_only)} PDL-only rows -> {PDL_CSV}", flush=True)

    # ---- summaries
    boards_meta = {}
    with open(BOARDS_CSV, encoding="utf-8") as f:
        for b in csv.DictReader(f):
            boards_meta[b["name"].strip()] = b
    state = {}
    if os.path.exists(STATE_JSON):
        with open(STATE_JSON, encoding="utf-8") as f:
            state = json.load(f)
    fetched_by_name = {v.get("name"): v.get("fetched", "")
                       for v in state.values()}

    board_stats = {}
    comp_stats = {}
    for r in rows:
        for bn in r[BOARDS]:
            st = board_stats.setdefault(bn, [0, set()])
            st[0] += 1
            st[1].add(r[IDX["company"]])
        ck = (r[IDX["company"]], r[ORG_SLUG])
        cs = comp_stats.setdefault(ck, [0, set(), r[IDX["company_stage"]],
                                        r[IDX["company_size"]],
                                        r[IDX["company_industry"]]])
        cs[0] += 1
        cs[1] |= r[BOARDS]

    # ---- workbook
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def headers(ws, names):
        out = []
        for n_ in names:
            c = WriteOnlyCell(ws, value=n_)
            c.font = bold
            out.append(c)
        return out

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    done_boards = sum(1 for v in state.values() if v.get("status") == "done")
    lines = [
        ("Getro job boards — full aggregation", ""),
        ("Snapshot (UTC)", SNAPSHOT),
        ("", ""),
        ("Boards fetched", done_boards or len(board_stats)),
        ("Raw job rows", raw_rows),
        ("Unique jobs (deduped)", len(rows)),
        ("Companies", len(comp_stats)),
        ("", ""),
        ("How deduping works", "Cross-posted jobs (same Getro job id, or same "
         "company + title + apply URL) are one row; every board carrying the "
         "job is listed in 'boards'."),
        ("Master CSV", "getro_jobs_master.csv holds the complete deduped "
         "table."),
    ]
    if split:
        lines += [
            ("Why this workbook is split", f"The deduped total ({len(rows)}) "
             "exceeds Excel's 1,048,576-row sheet limit. 'All Jobs' holds "
             f"every job listed on the {len(board_stats) - 1} curated "
             "ecosystem boards (including jobs cross-listed on People Data "
             f"Labs). The {len(pdl_only)} jobs found ONLY on the People Data "
             "Labs firehose board are in getro_jobs_pdl_only.csv."),
        ]
    lines += [
        ("", ""),
        ("Column", "Meaning"),
        ("title / company / location", "Role, employer, posting location(s)"),
        ("work_mode", "On-site / Hybrid / Remote"),
        ("seniority", "Getro's level tag (Entry Level, Associate, Mid "
         "Senior, Director, Internship, Senior)"),
        ("posted", "First-seen posting date (UTC)"),
        ("salary_min / salary_max", "Published compensation, whole units of "
         "salary_currency per salary_period; blank when not public"),
        ("equity", "Posting advertises equity"),
        ("skills", "Top skills tagged by Getro (max 8)"),
        ("company_stage / size / industry", "Employer funding stage, head "
         "count, top industry tags"),
        ("apply_url", "Direct application link"),
        ("board_job_url", "The posting on a Getro board (backup if the apply "
         "link dies)"),
        ("boards / board_count", "Every Getro network listing this job"),
        ("source", "Where Getro found it (career page, manual, ...)"),
        ("getro_job_id", "Stable id inside Getro"),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    widths = [46, 26, 30, 10, 12, 11, 11, 11, 9, 12, 7, 44, 13, 12, 34, 46,
              46, 40, 12, 12, 13]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    n_rows = min(len(eco), XLSX_ROW_CAP)
    ws.auto_filter.ref = f"A1:{get_column_letter(N)}{n_rows + 1}"
    ws.append(headers(ws, COLUMNS))
    for r in eco[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for col, wdt in zip("ABCDEF", (42, 44, 16, 12, 12, 12)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(board_stats) + 1}"
    ws.append(headers(ws, ["board", "url", "host_type", "job_rows",
                           "unique_jobs", "companies"]))
    for bn in sorted(board_stats, key=lambda x: -board_stats[x][0]):
        meta = boards_meta.get(bn, {})
        ws.append([safe(bn), meta.get("board_url", ""),
                   meta.get("host_type", ""), fetched_by_name.get(bn, ""),
                   board_stats[bn][0], len(board_stats[bn][1])])

    ws = wb.create_sheet("Companies")
    for col, wdt in zip("ABCDEF", (34, 14, 10, 44, 10, 10)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(comp_stats) + 1}"
    ws.append(headers(ws, ["company", "stage", "size", "industry", "jobs",
                           "boards"]))
    for (name, _slug), (n_, bset, stage, size, ind) in sorted(
            comp_stats.items(), key=lambda kv: -kv[1][0]):
        ws.append([safe(name), stage, size, safe(ind), n_, len(bset)])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
