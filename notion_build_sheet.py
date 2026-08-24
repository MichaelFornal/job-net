#!/usr/bin/env python3
"""
notion_build_sheet.py — notion_all_jobs.jsonl -> notion_jobs_master.csv
+ Notion_All_Jobs_<date>.xlsx.

Column names are the all_jobs_build.py names verbatim so the family needs
nothing more than `load_notion = _simple_loader("notion", "ats")` there.

Two things are unusual for this repo and both are deliberate:

  posted     REAL, not first-seen.  Every Notion row block carries
             created_time in epoch ms, so ~100% of rows land with
             precision="created".  The ledger still freezes it (see
             notion_fetch_jobs.py) because a date that moves between runs
             marches rows backwards past the all_jobs_build cutoff.
  location   free text, per-tenant.  "Flex-Boston", "Remote (US-timezones)",
             "Bengaluru (IISc Campus)".  city/state are parsed only when the
             shape is unambiguous; a wrong state is worse than a blank one
             because ma_remote_build.geo_of routes on it.

job_url is https://{host}.notion.site/{32hex}, the shape
all_jobs_build.canon_key turns into ats|notion||{id} — keyed on the row id
ALONE, like Dover, because the same posting is reachable under a bare slug and
under Title-<hex>.
"""
import csv
import json
import os
import re
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

import contract_common as cc

BASE = os.path.dirname(os.path.abspath(__file__))
SMOKE = int(os.environ.get("SMOKE") or 0)
PFX = "smoke_" if SMOKE else ""

IN_JSONL = os.path.join(BASE, PFX + "notion_all_jobs.jsonl")
BOARDS_CSV = os.path.join(BASE, PFX + "notion_job_boards.csv")
OUT_CSV = os.path.join(BASE, PFX + "notion_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, "%sNotion_All_Jobs_%s.xlsx" % (
    "Smoke_" if SMOKE else "", SNAPSHOT))
XLSX_ROW_CAP = 1_048_000

COLUMNS = ["title", "company", "location", "work_mode", "remote",
           "employment_type", "department", "city", "state", "country",
           "posted", "posted_precision", "apply_url", "job_url",
           "board", "board_url", "notion_job_id"]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)

REMOTE_RE = re.compile(r"\bremote|anywhere|work from home|wfh\b", re.I)
HYBRID_RE = re.compile(r"hybrid|flex-", re.I)
ONSITE_RE = re.compile(r"on-?site|in-?office|in-?person", re.I)

US_STATES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID",
    "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS",
    "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH", "OK",
    "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV",
    "WI", "WY", "DC",
}
# "Austin, TX" / "Conway, AR, USA"
STATE_RE = re.compile(r",\s*([A-Z]{2})(?:\s*,|\s*$)")
# Cities that pin a US state without one being written down.  Kept
# deliberately short: a wrong state routes a job into the wrong triage lane.
CITY_STATE = {
    "boston": "MA", "cambridge": "MA", "somerville": "MA",
    "new york": "NY", "brooklyn": "NY", "nyc": "NY",
    "san francisco": "CA", "sf": "CA", "los angeles": "CA",
    "palo alto": "CA", "mountain view": "CA", "oakland": "CA",
    "san jose": "CA", "santa monica": "CA",
    "seattle": "WA", "austin": "TX", "dallas": "TX", "houston": "TX",
    "chicago": "IL", "denver": "CO", "boulder": "CO", "atlanta": "GA",
    "miami": "FL", "washington": "DC", "philadelphia": "PA",
    "portland": "OR", "san diego": "CA", "providence": "RI",
}
NON_US_RX = re.compile(
    r"\b(london|berlin|paris|amsterdam|dublin|madrid|barcelona|lisbon|milan|"
    r"rome|zurich|munich|vienna|stockholm|copenhagen|oslo|helsinki|warsaw|"
    r"prague|budapest|athens|istanbul|dubai|tel aviv|bangalore|bengaluru|"
    r"mumbai|delhi|hyderabad|pune|chennai|singapore|tokyo|seoul|beijing|"
    r"shanghai|hong kong|sydney|melbourne|auckland|toronto|vancouver|"
    r"montreal|s[aã]o paulo|rio de janeiro|buenos aires|mexico city|"
    r"bogot[aá]|lima|santiago|lagos|nairobi|cairo|cape town|johannesburg|"
    r"uk|united kingdom|germany|france|spain|italy|india|brazil|canada|"
    r"australia|japan|china|nigeria|kenya|netherlands|poland|portugal)\b",
    re.I)

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def geo(loc):
    """(city, state, country) — only when unambiguous.

    A blank is safe; a wrong state is not.  ma_remote_build.geo_of routes the
    MA lane off `state`, so guessing here would put jobs in a lane that is
    actually applied from."""
    city = state = country = ""
    t = (loc or "").strip()
    if not t or REMOTE_RE.search(t):
        return city, state, country
    m = STATE_RE.search(t)
    if m and m.group(1) in US_STATES:
        state, country = m.group(1), "US"
        city = t.split(",")[0].strip()
        return I(city), I(state), I(country)
    if NON_US_RX.search(t):
        return "", "", ""
    head = re.split(r"[,(/|]", t)[0].strip().lower()
    head = re.sub(r"^(flex|hybrid|onsite|on-site)[\s-]+", "", head)
    if head in CITY_STATE:
        return I(head.title()), I(CITY_STATE[head]), I("US")
    return "", "", ""


def norm(rec):
    j = rec.get("job") or {}
    title = (j.get("title") or "").strip()
    loc = I((j.get("location") or "").strip())
    blob = " ".join((loc, j.get("remote_txt") or "", j.get("employment_type")
                     or "", title))

    if REMOTE_RE.search(blob):
        wm = "Remote"
    elif HYBRID_RE.search(blob):
        wm = "Hybrid"
    elif ONSITE_RE.search(blob) or loc:
        wm = "On-site"
    else:
        wm = ""

    city, state, country = geo(loc)
    url = (j.get("apply_url") or "").strip()
    jurl = (j.get("job_url") or "").strip()

    r = [""] * N
    r[IDX["title"]] = title
    r[IDX["company"]] = I((j.get("company") or "").strip())
    r[IDX["location"]] = loc
    r[IDX["work_mode"]] = I(wm)
    r[IDX["remote"]] = "True" if wm == "Remote" else ""
    r[IDX["employment_type"]] = I(cc.classify_any(j.get("employment_type"),
                                                  title))
    r[IDX["department"]] = I((j.get("department") or "").strip())
    r[IDX["city"]] = city
    r[IDX["state"]] = state
    r[IDX["country"]] = country
    r[IDX["posted"]] = I(j.get("posted") or "")
    r[IDX["posted_precision"]] = I(j.get("posted_precision") or "")
    r[IDX["apply_url"]] = url or jurl
    r[IDX["job_url"]] = jurl or url
    r[IDX["board"]] = I(rec.get("host") or "")
    r[IDX["board_url"]] = I(rec.get("board_url") or "")
    r[IDX["notion_job_id"]] = j.get("id") or ""
    return r


def main():
    if not os.path.exists(IN_JSONL):
        raise SystemExit("missing %s — run notion_fetch_jobs.py" % IN_JSONL)

    by_id, raw, tiers = {}, 0, {}
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            raw += 1
            r = norm(rec)
            if not r[IDX["title"]] or not r[IDX["company"]]:
                continue
            tiers[(rec.get("job") or {}).get("tier") or "?"] = \
                tiers.get((rec.get("job") or {}).get("tier") or "?", 0) + 1
            key = r[IDX["notion_job_id"]] or "%s|%s|%s" % (
                r[IDX["company"]], r[IDX["title"]], r[IDX["location"]])
            by_id.setdefault(key, r)
    n_ids = len(by_id)

    by_sig = {}
    for r in by_id.values():
        sig = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               (r[IDX["apply_url"]] or r[IDX["location"]]).lower())
        by_sig.setdefault(sig, r)
    rows = list(by_sig.values())
    rows.sort(key=lambda r: (r[IDX["posted"]] or "", r[IDX["company"]]),
              reverse=True)
    print("raw rows %d -> unique ids %d -> deduped %d  (tiers %s)"
          % (raw, n_ids, len(rows), tiers), flush=True)

    n_real = sum(1 for r in rows if r[IDX["posted_precision"]] == "created")
    pct = 100.0 * n_real / max(1, len(rows))
    if pct < 60:
        print("  WARNING: only %.0f%% of rows carry a real created date — "
              "expected ~100%%; check that queryCollection is still returning "
              "created_time on row blocks" % pct, flush=True)
    else:
        print("  %.0f%% of rows carry a real created date" % pct, flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print("wrote %s" % OUT_CSV, flush=True)
    write_xlsx(rows, tiers)


def write_xlsx(rows, tiers):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    n_rem = sum(1 for r in rows if r[IDX["remote"]])
    for a, b in [
        ("Notion — all jobs", "Snapshot %s" % SNAPSHOT),
        ("Rows", "%s deduplicated postings from %s boards"
                 % (format(len(rows), ","),
                    format(len({r[IDX["board"]] for r in rows}), ","))),
        ("Remote", "%s inferred remote from free-text location"
                   % format(n_rem, ",")),
        ("Source", "Public notion.site careers pages via the unauthenticated "
                   "/api/v3 endpoints: loadPageChunk -> queryCollection. No "
                   "key, no cookie. www.notion.so and <sub>.notion.site "
                   "return identical responses."),
        ("Tiers", "db = rows of a Notion database (%d). pages = child "
                  "sub-pages of a careers page (%d). Prose careers pages are "
                  "NOT parsed: a header can be a job title on one board and a "
                  "responsibility on the next, so parsing them would "
                  "fabricate rows."
                  % (tiers.get("db", 0), tiers.get("pages", 0))),
        ("posted", "REAL dates, not first-seen: every Notion row block "
                   "carries created_time. Frozen in notion_posted_ledger.json "
                   "on first sight so a row cannot march backwards past the "
                   "all_jobs_build cutoff while still live."),
        ("Not a job board", "A Notion database is not automatically a job "
                            "board. reachcapital.notion.site — 340 rows in "
                            "all_jobs_master.csv — is a VC READING LIST "
                            "(\"Our Reading Syllabus\", title property "
                            "\"Publication\"). Boards must clear a gate on "
                            "collection name, schema, and a >=30% role-word "
                            "rate across row titles."),
        ("Company", "spaceName is used only when it corroborates the host; "
                    "aliendesign.notion.site reports \"Yash repo\". "
                    "Multi-employer boards emit company per-row or are "
                    "dropped — never the workspace name."),
        ("Dedupe", "job_url is {host}.notion.site/{32hex}; "
                   "all_jobs_build.canon_key keys on the row id ALONE, "
                   "because one posting is reachable both as a bare slug and "
                   "as Title-<hex>."),
    ]:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    ws.freeze_panes = "A2"
    widths = {"title": 52, "company": 26, "location": 32, "work_mode": 10,
              "remote": 8, "employment_type": 14, "department": 22,
              "city": 18, "state": 7, "country": 9, "posted": 11,
              "posted_precision": 15, "apply_url": 56, "job_url": 56,
              "board": 24, "board_url": 46, "notion_job_id": 34}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    hdr = []
    for c in COLUMNS:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(N),
                                      min(len(rows), XLSX_ROW_CAP) + 1)
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABCDE", (28, 30, 10, 12, 46)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    hdr = []
    for c in ("board", "company", "jobs", "newest_post", "board_url"):
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    stats = {}
    for r in rows:
        st = stats.setdefault(r[IDX["board"]],
                              [r[IDX["company"]], 0, "", r[IDX["board_url"]]])
        st[1] += 1
        if (r[IDX["posted"]] or "") > st[2]:
            st[2] = r[IDX["posted"]] or ""
    for b, (comp, n, newest, url) in sorted(stats.items(),
                                            key=lambda kv: -kv[1][1]):
        ws.append([safe(b), safe(comp), n, newest, safe(url)])

    wb.save(OUT_XLSX)
    print("wrote %s" % OUT_XLSX, flush=True)


if __name__ == "__main__":
    sys.exit(main())
