#!/usr/bin/env python3
"""
dover_build_sheet.py — dover_all_jobs.jsonl -> dover_jobs_master.csv +
Dover_All_Jobs_<date>.xlsx.

apply_url = job_url = https://app.dover.com/apply/{slug}/{job_id}, the shape
already present in all_jobs_master.csv (292 such URLs today).  Note that Dover
serves the SAME posting under two first-segment forms — the company slug
(app.dover.com/apply/transfyr/{job}) and the careers-page UUID
(app.dover.com/apply/dfdb6661-.../{job}) — which is why
all_jobs_build.canon_key keys Dover on the job UUID alone and ignores that
segment.  Emitting the slug form here is purely for humans.

Unlike Rippling/Personio/YC this family has REAL dates: the detail endpoint's
`created` is a true ISO timestamp, so posted_precision is "created" and the
ledger is only a fallback for jobs whose detail call failed.
"""
import csv
import json
import os
import re
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

BASE = os.path.dirname(os.path.abspath(__file__))
SMOKE = int(os.environ.get("SMOKE") or 0)
PFX = "smoke_" if SMOKE else ""

IN_JSONL = os.path.join(BASE, PFX + "dover_all_jobs.jsonl")
LEDGER_JSON = os.path.join(BASE, PFX + "dover_posted_ledger.json")
OUT_CSV = os.path.join(BASE, PFX + "dover_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, "%sDover_All_Jobs_%s.xlsx" % (
    "Smoke_" if SMOKE else "", SNAPSHOT))
XLSX_ROW_CAP = 1_048_000

COLUMNS = ["title", "company", "company_domain", "location", "work_mode",
           "remote", "employment_type", "city", "state", "country",
           "secondary_locations", "posted", "posted_precision",
           "salary_min", "salary_max", "salary_currency", "salary_period",
           "equity_min", "equity_max", "visa_support",
           "apply_url", "job_url", "board", "board_url", "dover_job_id"]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)

WM_MAP = {"ONSITE": "On-site", "IN_OFFICE": "On-site", "REMOTE": "Remote",
          "HYBRID": "Hybrid"}
ET_MAP = {"FULL_TIME": "Full-time", "PART_TIME": "Part-time",
          "CONTRACT": "Contract", "INTERNSHIP": "Internship",
          "TEMPORARY": "Temp"}
PERIOD_MAP = {"YEARLY": "year", "MONTHLY": "month", "WEEKLY": "week",
              "DAILY": "day", "HOURLY": "hour"}
REMOTE_RE = re.compile(r"remote|anywhere", re.I)

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def loc_parts(j):
    """-> (display, city, state, country, secondary)."""
    locs = j.get("locations")
    if not isinstance(locs, list) or not locs:
        one = j.get("location")
        if isinstance(one, dict):
            locs = [one]
        elif isinstance(one, str) and one:
            return one, "", "", "", ""
        else:
            return "", "", "", "", ""
    prim = None
    for x in locs:
        if isinstance(x, dict) and x.get("is_primary"):
            prim = x
            break
    if prim is None:
        prim = next((x for x in locs if isinstance(x, dict)), {})
    opt = prim.get("location_option") or prim
    disp = (opt.get("display_name") or opt.get("name") or "").strip()
    city = (opt.get("city") or "").strip()
    state = (opt.get("state") or "").strip()
    country = (opt.get("country") or "").strip()
    others = []
    for x in locs:
        if x is prim or not isinstance(x, dict):
            continue
        o = x.get("location_option") or x
        d = (o.get("display_name") or o.get("name") or "").strip()
        if d and d != disp:
            others.append(d)
    return disp, city, state, country, "; ".join(others)


def norm(rec, ledger):
    j = rec.get("job") or {}
    jid = str(j.get("id") or "").strip()
    slug = rec.get("slug") or ""
    title = (j.get("title") or "").strip()
    disp, city, state, country, others = loc_parts(j)

    wt = (j.get("workplace_type") or "").strip().upper()
    wm = WM_MAP.get(wt, "")
    if not wm:
        wm = "Remote" if REMOTE_RE.search(disp) else ("On-site" if disp else "")

    comp = j.get("compensation") or {}
    lo = comp.get("lower_bound")
    hi = comp.get("upper_bound")
    has = lo is not None or hi is not None
    cur = (comp.get("currency_code") or "") if has else ""
    per = PERIOD_MAP.get((comp.get("salary_range_type") or "").upper(),
                         "year" if has else "")
    et = ET_MAP.get((comp.get("employment_type") or "").upper(), "")

    entry = ledger.get(jid) or {}
    created = (j.get("created") or "")[:10]
    posted = created or entry.get("posted") or ""
    prec = "created" if created else (entry.get("precision") or "")

    url = ("https://app.dover.com/apply/%s/%s" % (slug or rec.get("uuid"), jid)
           if jid else "")

    r = [""] * N
    r[IDX["title"]] = title
    r[IDX["company"]] = I(j.get("client_name") or rec.get("company") or slug)
    r[IDX["company_domain"]] = I(j.get("client_domain") or "")
    r[IDX["location"]] = I(disp)
    r[IDX["work_mode"]] = I(wm)
    r[IDX["remote"]] = "True" if wm == "Remote" else ""
    r[IDX["employment_type"]] = I(et)
    r[IDX["city"]] = I(city)
    r[IDX["state"]] = I(state)
    r[IDX["country"]] = I(country)
    r[IDX["secondary_locations"]] = I(others)
    r[IDX["posted"]] = I(posted)
    r[IDX["posted_precision"]] = I(prec)
    r[IDX["salary_min"]] = str(int(lo)) if isinstance(lo, (int, float)) else ""
    r[IDX["salary_max"]] = str(int(hi)) if isinstance(hi, (int, float)) else ""
    r[IDX["salary_currency"]] = I(cur)
    r[IDX["salary_period"]] = I(per)
    eq_lo = comp.get("equity_lower_bound")
    eq_hi = comp.get("equity_upper_bound")
    r[IDX["equity_min"]] = str(eq_lo) if eq_lo is not None else ""
    r[IDX["equity_max"]] = str(eq_hi) if eq_hi is not None else ""
    r[IDX["visa_support"]] = I(str(j.get("visa_support") or ""))
    r[IDX["apply_url"]] = url
    r[IDX["job_url"]] = url
    r[IDX["board"]] = I(slug or rec.get("uuid") or "")
    r[IDX["board_url"]] = I(rec.get("board_url") or "")
    r[IDX["dover_job_id"]] = jid
    return r


def main():
    if not os.path.exists(IN_JSONL):
        raise SystemExit("missing %s — run dover_fetch_jobs.py" % IN_JSONL)
    ledger = {}
    if os.path.exists(LEDGER_JSON):
        try:
            with open(LEDGER_JSON, encoding="utf-8") as f:
                ledger = json.load(f)
        except (OSError, json.JSONDecodeError):
            ledger = {}

    by_id, raw = {}, 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            raw += 1
            r = norm(rec, ledger)
            if not r[IDX["title"]]:
                continue
            key = r[IDX["dover_job_id"]] or "%s|%s" % (r[IDX["company"]],
                                                       r[IDX["title"]])
            by_id.setdefault(key, r)
    n_ids = len(by_id)

    by_sig = {}
    for r in by_id.values():
        sig = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               (r[IDX["apply_url"]] or r[IDX["location"]]).lower())
        by_sig.setdefault(sig, r)
    rows = list(by_sig.values())
    rows.sort(key=lambda r: (r[IDX["posted"]] or "", r[IDX["company"]]),
              reverse=True)
    print("raw rows %d -> unique ids %d -> deduped %d"
          % (raw, n_ids, len(rows)), flush=True)

    n_created = sum(1 for r in rows
                    if r[IDX["posted_precision"]] == "created")
    if rows and n_created / len(rows) < 0.8:
        print("  WARNING: only %d/%d rows carry a real created date — the "
              "detail endpoint may be degraded" % (n_created, len(rows)),
              flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print("wrote %s" % OUT_CSV, flush=True)
    write_xlsx(rows)


def write_xlsx(rows):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    n_sal = sum(1 for r in rows if r[IDX["salary_min"]])
    for a, b in [
        ("Dover — all jobs", "Snapshot %s" % SNAPSHOT),
        ("Rows", "%s deduplicated postings" % format(len(rows), ",")),
        ("With salary", "%s carry structured compensation bounds"
         % format(n_sal, ",")),
        ("Why it matters", "Dover sells almost exclusively to pre-seed to "
                           "Series-B startups, which is the segment the rest "
                           "of the corpus is thinnest on."),
        ("Source", "app.dover.com/api/v1/careers-page/{uuid}/jobs for the "
                   "list, then /api/v1/inbound/application-portal-job/{id} "
                   "per job for the date, salary and company."),
        ("Caution", "app.dover.com returns HTTP 200 with a 3,988-byte HTML "
                    "SPA shell on ANY unrecognised path, so every response "
                    "here is content-type checked rather than status checked."),
        ("Dedupe", "The same posting is reachable under both "
                   "/apply/{slug}/{job} and /apply/{careers-uuid}/{job}, so "
                   "canon_key keys Dover on the globally-unique job UUID "
                   "alone and ignores that segment."),
    ]:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    ws.freeze_panes = "A2"
    widths = {"title": 50, "company": 26, "company_domain": 24,
              "location": 30, "work_mode": 10, "remote": 8,
              "employment_type": 14, "city": 18, "state": 16, "country": 9,
              "secondary_locations": 30, "posted": 11, "posted_precision": 15,
              "salary_min": 11, "salary_max": 11, "salary_currency": 9,
              "salary_period": 9, "equity_min": 10, "equity_max": 10,
              "visa_support": 12, "apply_url": 56, "job_url": 56,
              "board": 24, "board_url": 46, "dover_job_id": 38}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    hdr = []
    for c in COLUMNS:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(N),
                                      min(len(rows), XLSX_ROW_CAP) + 1)
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABCD", (26, 26, 10, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    hdr = []
    for c in ("board", "company", "jobs", "newest_post"):
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    stats = {}
    for r in rows:
        st = stats.setdefault(r[IDX["board"]], [r[IDX["company"]], 0, ""])
        st[1] += 1
        if (r[IDX["posted"]] or "") > st[2]:
            st[2] = r[IDX["posted"]] or ""
    for b, (comp, n, newest) in sorted(stats.items(),
                                       key=lambda kv: -kv[1][1]):
        ws.append([safe(b), safe(comp), n, newest])

    wb.save(OUT_XLSX)
    print("wrote %s" % OUT_XLSX, flush=True)


if __name__ == "__main__":
    sys.exit(main())
