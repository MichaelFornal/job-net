#!/usr/bin/env python3
"""
contract_build.py — merge every contract_norm/*.jsonl into one deduplicated
master and an organized workbook.

Dedupe, two passes:
  1. exact (source, job_id) — within-source repeats.
  2. cross-source: (company, title, apply_url) or (company, title, location)
     when no apply_url — the same gig posted on a company ATS board AND
     surfaced on a VC/remote aggregator collapses to one row, its `sources`
     and `boards` accumulated, newest posted date kept.

Outputs:
  contract_master.csv              full deduped table
  Contract_Gigs_<date>.xlsx        README / All Gigs / Summary / Sources /
                                   Boards
"""
import csv
import json
import os
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

import contract_common as cc
import build_sources

BASE = os.path.dirname(os.path.abspath(__file__))
SNAP = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_CSV = os.path.join(BASE, "contract_master.csv")
OUT_XLSX = os.path.join(BASE, f"Contract_Gigs_{SNAP}.xlsx")
XLSX_CAP = 1_048_000

COLS = cc.COLUMNS + ["sources", "n_sources"]


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def merge(a, b):
    a["_sources"] |= b["_sources"]
    a["_boards"] |= b["_boards"]
    if b["posted"] > a["posted"]:
        a["posted"] = b["posted"]
    for k in ("apply_url", "location", "city", "state", "country",
              "salary_min", "salary_max", "salary_currency", "work_mode",
              "company"):
        if not a.get(k) and b.get(k):
            a[k] = b[k]


def main():
    by_id = {}
    dup_id = 0
    raw = 0
    dropped_ft = 0
    # One generation per source (cc.norm_files), then a PER-ROW type test.
    #
    # This used to skip any basename containing "_all", on the reasoning that
    # those were the KEEP_ALL dumps and everything else was contingent-only.
    # That inverted on 2026-07-29: ~25 new pipes (workday, careerjet, oracle,
    # icims, jobgether, themuse, intlmisc, intlboards, adzuna, adp, ...) write
    # full-employment-type dumps under names with no "_all" in them.  The
    # filename rule then admitted ~700k Full-time rows into a master documented
    # as contingent-only, while still discarding the KEEP_ALL files that hold
    # the freshest contract data.  classify() asks the row instead, so a new
    # pipe cannot re-introduce the bug by choosing a filename.
    for path in cc.norm_files():
        with open(path, encoding="utf-8") as f:
            for line in f:
                raw += 1
                try:
                    r = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if not (r.get("title") and str(r.get("title")).strip()):
                    continue
                label = cc.classify(r.get("employment_type"), r.get("title"))
                if not label:
                    dropped_ft += 1
                    continue
                r["employment_type"] = label
                r["_sources"] = {r.get("source")}
                r["_boards"] = {f"{r.get('source')}:{r.get('board')}"}
                jid = r.get("job_id")
                k = (r.get("source"), str(jid)) if jid else None
                if k and k in by_id:
                    merge(by_id[k], r)
                    dup_id += 1
                elif k:
                    by_id[k] = r
                else:
                    by_id[id(r)] = r
    after_id = len(by_id)

    by_key = {}
    for r in by_id.values():
        comp = (r.get("company") or "").strip().lower()
        title = (r.get("title") or "").strip().lower()
        au = (r.get("apply_url") or "").strip().lower()
        k = (comp, title, au or "loc:" + (r.get("location") or "").lower())
        if k in by_key:
            merge(by_key[k], r)
        else:
            by_key[k] = r
    rows = list(by_key.values())
    for r in rows:
        srcs = sorted(x for x in r["_sources"] if x)
        r["sources"] = ", ".join(srcs)
        r["n_sources"] = len(srcs)
    rows.sort(key=lambda r: (r.get("posted") or "", r.get("n_sources", 0)),
              reverse=True)
    print(f"raw {raw} -> contingent {raw - dropped_ft} "
          f"({dropped_ft} non-contingent dropped) "
          f"-> after id-dedupe {after_id} ({dup_id} id dups) "
          f"-> deduped {len(rows)}", flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for r in rows:
            w.writerow([r.get(c, "") for c in COLS])
    print(f"wrote {OUT_CSV}", flush=True)

    # ---- aggregates
    by_source = {}
    by_type = {}
    by_board = {}
    for r in rows:
        s = r.get("source") or ""
        by_source[s] = by_source.get(s, 0) + 1
        et = r.get("employment_type") or "(none)"
        by_type[et] = by_type.get(et, 0) + 1
        bk = (s, r.get("board") or "")
        st = by_board.setdefault(bk, [0, r.get("company") or "",
                                       r.get("board_url") or "", ""])
        st[0] += 1
        if (r.get("posted") or "") > st[3]:
            st[3] = r.get("posted") or ""

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def hdr(ws, names):
        out = []
        for n in names:
            c = WriteOnlyCell(ws, value=n)
            c.font = bold
            out.append(c)
        return out

    # README
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 116
    n_remote = sum(1 for r in rows if r.get("remote"))
    lines = [
        ("Contract & temp gigs — master aggregation", ""),
        ("Snapshot (UTC)", SNAP),
        ("", ""),
        ("Unique gigs (deduped)", len(rows)),
        ("Raw rows harvested", raw),
        ("Remote gigs", n_remote),
        ("Distinct sources", len(by_source)),
        ("", ""),
        ("What this is", "Every contract / temp / freelance / fractional / "
         "interim role we could reach through an open, unauthenticated feed, "
         "across staffing ATSs, freelance marketplaces, remote boards, VC "
         "portfolio boards, and first-party company ATS boards. Full source "
         "directory (incl. gated networks that need login/keys) is on the "
         "Sources tab."),
        ("How a gig qualifies", "Either the source's own employment-type "
         "field says Contract/Temp/Freelance/Seasonal/Interim, or (for feeds "
         "with no type field: Getro, some RSS) the title does. Permanent-"
         "contract wording (CDI, 'permanent contract', unbefristet) is "
         "excluded."),
        ("Dedupe", "Same posting collapses across sources on (company, "
         "title, apply URL) or (company, title, location). The 'sources' and "
         "'n_sources' columns show where a gig was seen; n_sources>1 is a "
         "cross-confirmed listing."),
        ("Master CSV", "contract_master.csv holds the complete table."),
        ("", ""),
        ("Column", "Meaning"),
        ("source / source_type", "Feed it came from and its family (staffing"
         "_ats, marketplace, remote_board, vc_board, ats, agency)."),
        ("employment_type", "Normalized contingent label: Contract, Temp, "
         "Freelance, Contract-to-Hire, Interim, Seasonal, Per Diem, "
         "Fractional."),
        ("work_mode / remote", "Remote / Hybrid / On-site, and a remote "
         "boolean."),
        ("location / city / state / country", "Where the gig sits (freeform "
         "+ parsed parts where the source gave them)."),
        ("salary_min/max/currency/period", "Pay when the source publishes it "
         "(hourly for many contracts; period says Hour/Year/Project)."),
        ("posted", "Publish date (UTC) where available."),
        ("apply_url", "Direct link to apply / the posting. Bullhorn rows have "
         "no public link (agency-portal only)."),
        ("board / board_url", "The specific account/board on the source."),
        ("sources / n_sources", "All feeds this gig appeared on; count."),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    # All Gigs
    ws = wb.create_sheet("All Gigs")
    widths = [12, 13, 26, 44, 15, 9, 30, 16, 12, 9, 7, 10, 10, 8, 8, 11, 46,
              20, 18, 34, 9]
    for i, wd in enumerate(widths[:len(COLS)], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    n_x = min(len(rows), XLSX_CAP)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n_x + 1}"
    ws.append(hdr(ws, COLS))
    for r in rows[:XLSX_CAP]:
        ws.append([safe(r.get(c, "")) for c in COLS])
    if len(rows) > XLSX_CAP:
        print(f"NOTE: {len(rows) - XLSX_CAP} rows beyond Excel cap (CSV only)",
              flush=True)

    # Summary
    ws = wb.create_sheet("Summary")
    for col, wd in zip("ABCD", (18, 16, 12, 12)):
        ws.column_dimensions[col].width = wd
    ws.append(hdr(ws, ["by source", "gigs", "", ""]))
    for s, n in sorted(by_source.items(), key=lambda kv: -kv[1]):
        ws.append([s, n])
    ws.append(["", "", "", ""])
    ws.append(hdr(ws, ["by employment_type", "gigs", "", ""]))
    for t, n in sorted(by_type.items(), key=lambda kv: -kv[1]):
        ws.append([t, n])
    ws.append(["", "", "", ""])
    ws.append([WriteOnlyCell(ws, value="TOTAL"), len(rows)])

    # Sources
    ws = wb.create_sheet("Sources")
    scols = build_sources.COLS
    for col, wd in zip("ABCDEFG", (30, 16, 15, 34, 12, 52, 90)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    src_rows = build_sources.rows()
    ws.auto_filter.ref = f"A1:G{len(src_rows) + 1}"
    ws.append(hdr(ws, scols))
    for sr in src_rows:
        c_last = WriteOnlyCell(ws, value=sr["notes"])
        c_last.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([sr["source"], sr["category"], sr["access"],
                   sr["contract_signal"], sr["status"], sr["endpoint"],
                   c_last])

    # Boards
    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABCDE", (14, 30, 46, 9, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(by_board) + 1}"
    ws.append(hdr(ws, ["source", "board", "board_url", "gigs", "newest"]))
    for (s, b), st in sorted(by_board.items(), key=lambda kv: -kv[1][0]):
        ws.append([s, b, st[2], st[0], st[3]])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
