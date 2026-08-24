#!/usr/bin/env python3
"""Build Remote_Jobs_<date>.xlsx from the extracted remote CSV.

Tabs:
  Summary     counts + the exact rules used
  Remote Jobs every remote row in the master (is_remote from all_jobs_build)
  US Remote   the subset reachable from the US (see remote_us_classify)

Reads Remote_Jobs_<date>.csv (produced by remote_jobs_build.py) so the
1.15GB master does not have to be re-scanned.
"""
import csv
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from remote_us_classify import us_basis

HERE = os.path.dirname(os.path.abspath(__file__))
DATE = sys.argv[1] if len(sys.argv) > 1 else "2026-07-29"
SRC_CSV = os.path.join(HERE, f"Remote_Jobs_{DATE}.csv")
OUT_XLSX = os.path.join(HERE, f"Remote_Jobs_{DATE}.xlsx")
OUT_US_CSV = os.path.join(HERE, f"US_Remote_Jobs_{DATE}.csv")

XLSX_CAP = 1_048_000
US_KEEP = {"country", "state", "location", "worldwide-open"}

WIDTHS = [26, 44, 14, 12, 18, 9, 7, 13, 30, 14, 10, 12, 14, 11, 10, 10, 8, 9,
          7, 40, 13, 11, 30, 46, 46, 36, 8, 13, 12, 22, 9, 42]

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def main():
    csv.field_size_limit(10 ** 9)
    with open(SRC_CSV, newline="", encoding="utf-8") as f:
        rd = csv.reader(f)
        base = next(rd)
        i = {n: k for k, n in enumerate(base)}
        ci, si, li = i["country"], i["state"], i["location"]
        pos = i["country"] + 1
        cols = base[:pos] + ["us_basis"] + base[pos:]

        rows, us_rows = [], []
        basis_c, us_wm, us_src, us_state = Counter(), Counter(), Counter(), \
            Counter()
        for row in rd:
            b = I(us_basis(row[ci], row[si], row[li]))
            out = row[:pos] + [b] + row[pos:]
            rows.append(out)
            basis_c[b or "(not US)"] += 1
            if b in US_KEEP:
                us_rows.append(out)
                us_wm[row[i["work_mode"]] or "(not stated)"] += 1
                us_src[row[i["source"]] or "(none)"] += 1
                us_state[row[si] or "(not stated)"] += 1

    print(f"remote rows {len(rows):,} -> US remote {len(us_rows):,}",
          flush=True)

    with open(OUT_US_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(us_rows)
    print(f"wrote {OUT_US_CSV}", flush=True)

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def hdr(ws, names):
        out = []
        for nm in names:
            c = WriteOnlyCell(ws, value=nm)
            c.font = bold
            out.append(c)
        return out

    # ---- Summary
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 46
    ws.append(hdr(ws, ["Remote Jobs", DATE]))
    ws.append(["source", f"All_Jobs_{DATE} / all_jobs_master.csv"])
    ws.append(["master rows scanned", 2392615])
    ws.append(["remote rows (tab: Remote Jobs)", len(rows)])
    ws.append(["US remote rows (tab: US Remote)", len(us_rows)])
    ws.append([])
    ws.append(hdr(ws, ["us_basis", "jobs"]))
    for k, v in basis_c.most_common():
        ws.append([k, v])
    ws.append([])
    ws.append(hdr(ws, ["US Remote: work_mode", "jobs"]))
    for k, v in us_wm.most_common():
        ws.append([k, v])
    ws.append([])
    ws.append(hdr(ws, ["US Remote: state", "jobs"]))
    for k, v in us_state.most_common(25):
        ws.append([k, v])
    ws.append([])
    ws.append(hdr(ws, ["US Remote: source", "jobs"]))
    for k, v in us_src.most_common(30):
        ws.append([k, v])
    ws.append([])
    ws.append(hdr(ws, ["rules", ""]))
    ws.append(["remote (kept on 'Remote Jobs')",
               "ANY of: remote==True / work_mode==Remote / 'remote' in "
               "location -- the is_remote() rule all_jobs_build.py uses"])
    ws.append(["remote_signal column", "which of those three fired"])
    ws.append(["us_basis = country", "country field says US/USA/United States"])
    ws.append(["us_basis = state", "country blank, state is a US state"])
    ws.append(["us_basis = location",
               "location names the US or a US state (a non-US country only "
               "yields to explicit 'United States'/'USA'/'US Remote' wording)"])
    ws.append(["us_basis = worldwide-open",
               "no country, location says anywhere/worldwide/global -- open "
               "to US applicants; filter this out for US-specific only"])
    ws.append(["", "ISO-2 codes mean CA=Canada, DE=Germany, IN=India, so "
                   "2-letter tokens are read as US states only from the "
                   "state field or a ', XX' location segment"])

    # ---- data tabs
    def data_tab(base_name, table):
        n = len(cols)
        chunks = max(1, -(-len(table) // XLSX_CAP))
        for ci_ in range(chunks):
            name = base_name if ci_ == 0 else f"{base_name} {ci_ + 1}"
            w = wb.create_sheet(name)
            for j, wd in enumerate(WIDTHS[:n], 1):
                w.column_dimensions[get_column_letter(j)].width = wd
            w.freeze_panes = "A2"
            chunk = table[ci_ * XLSX_CAP:(ci_ + 1) * XLSX_CAP]
            w.auto_filter.ref = (f"A1:{get_column_letter(n)}"
                                 f"{len(chunk) + 1}")
            w.append(hdr(w, cols))
            for r in chunk:
                w.append([safe(v) for v in r])
            print(f"  tab '{name}': {len(chunk):,} rows", flush=True)

    data_tab("US Remote", us_rows)
    data_tab("Remote Jobs", rows)

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
