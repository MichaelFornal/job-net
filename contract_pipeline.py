#!/usr/bin/env python3
"""
contract_pipeline.py — personal pipeline workbook from contract_master.csv,
same shape as getro_build_pipeline.py / lever_build_pipeline.py:

  - every Remote gig
  - every On-site / Hybrid / untagged gig located in Massachusetts, Rhode
    Island, or the South Florida metros (Miami-Dade / Broward / Palm Beach),
    plus rows tagged only "Florida" statewide, labeled "FL (statewide tag)".

Flag + notes columns with an Interested / Maybe / No dropdown; Interested rows
turn green. Output: Contract_Pipeline_<date>.xlsx (Pipeline / Summary /
README).
"""
import csv
import os
import re
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
IN_CSV = os.path.join(BASE, "contract_master.csv")
SNAP = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Contract_Pipeline_{SNAP}.xlsx")
FLAG_CHOICES = "Interested,Maybe,No"

SFL_BARE = [
    "miami", "fort lauderdale", "boca raton", "coral gables", "delray beach",
    "boynton beach", "pompano beach", "deerfield beach", "hialeah", "doral",
    "aventura", "coral springs", "pembroke pines", "hallandale beach",
    "dania beach", "lauderhill", "tamarac", "opa-locka", "cutler bay",
    "palmetto bay", "sunny isles", "key biscayne", "wilton manors",
    "palm beach", "coconut creek", "cooper city", "juno beach", "tequesta",
    "miami-dade", "broward", "south florida",
]
SFL_NEED_FL = [
    "hollywood", "plantation", "sunrise", "davie", "weston", "miramar",
    "margate", "parkland", "homestead", "kendall", "jupiter", "wellington",
    "lake worth", "greenacres", "lantana", "riviera beach", "oakland park",
]
SFL_BARE_RE = re.compile(r"\b(" + "|".join(re.escape(c) for c in SFL_BARE)
                         + r")\b")
SFL_FL_RE = re.compile(r"\b(" + "|".join(re.escape(c) for c in SFL_NEED_FL)
                       + r"),?\s+(fl|florida)\b")
MA_RE = re.compile(r"massachusetts|,\s*ma\b|\bboston\b|greater boston")
RI_RE = re.compile(r"rhode island|,\s*ri\b|\bprovidence\b")
REGION_RANK = {"MA": 1, "RI": 2, "South FL": 3, "FL (statewide tag)": 4}


def statewide_fl(l):
    for seg in l.split(";"):
        first = seg.split(",", 1)[0].strip()
        if first in ("florida", "fl") or first.endswith("- florida") \
                or first.endswith("- fl"):
            return True
    return False


def regions_of(loc, state, country):
    if country and country.upper() not in ("US", "USA", "UNITED STATES", ""):
        return []
    l = (loc + " " + (state or "")).lower()
    out = []
    st = (state or "").strip().lower()
    if MA_RE.search(l) or st in ("ma", "massachusetts"):
        out.append("MA")
    if RI_RE.search(l) or st in ("ri", "rhode island"):
        out.append("RI")
    if SFL_BARE_RE.search(l) or SFL_FL_RE.search(l):
        out.append("South FL")
    elif "florida" in l and statewide_fl(l):
        out.append("FL (statewide tag)")
    return out


def is_remote(remote_val, work_mode, loc):
    if str(remote_val).strip().lower() in ("true", "1", "yes"):
        return True
    if (work_mode or "").strip().lower() == "remote":
        return True
    return "remote" in (loc or "").lower()


def main():
    csv.field_size_limit(10 ** 7)
    with open(IN_CSV, newline="", encoding="utf-8") as f:
        r = csv.reader(f)
        cols = next(r)
        i = {c: n for n, c in enumerate(cols)}
        rows = []
        n_remote = n_local = 0
        for row in r:
            loc = row[i["location"]]
            state = row[i.get("state", i["location"])] if "state" in i else ""
            country = row[i["country"]]
            regs = regions_of(loc, state, country)
            if is_remote(row[i["remote"]], row[i["work_mode"]], loc):
                rows.append((0, REGION_RANK.get(regs[0], 9) if regs else 9,
                             row[i["posted"]], "Remote",
                             "; ".join(regs), row))
                n_remote += 1
            elif regs:
                rows.append((1, REGION_RANK.get(regs[0], 9),
                             row[i["posted"]], "Local",
                             "; ".join(regs), row))
                n_local += 1
    rows.sort(key=lambda t: t[2], reverse=True)
    rows.sort(key=lambda t: (t[0], t[1]))
    print(f"pipeline rows: {len(rows)} ({n_remote} remote, {n_local} local)",
          flush=True)

    out_cols = ["flag", "notes", "bucket", "region"] + cols
    ncols = len(out_cols)
    nrows = len(rows)

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("Pipeline")
    widths = [11, 24, 9, 18] + [12, 13, 26, 44, 15, 9, 30, 16, 12, 9, 7, 10,
                                10, 8, 8, 11, 46, 20, 18, 34, 9]
    for idx, wd in enumerate(widths[:ncols], 1):
        ws.column_dimensions[get_column_letter(idx)].width = wd
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                        fill_type="solid")
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(ncols)}{nrows + 1}",
        FormulaRule(formula=['$A2="Interested"'], fill=green,
                    stopIfTrue=False))
    hdr = []
    for c in out_cols:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    for _, _, _, bucket, region, row in rows:
        ws.append(["", "", bucket, region] + row)

    ws = wb.create_sheet("Summary")
    counts = {}
    src_counts = {}
    for _, _, _, bucket, region, row in rows:
        counts[(bucket, region or "(anywhere)")] = counts.get(
            (bucket, region or "(anywhere)"), 0) + 1
        src_counts[row[0]] = src_counts.get(row[0], 0) + 1
    for col, wd in zip("ABC", (12, 26, 12)):
        ws.column_dimensions[col].width = wd
    ws.append([WriteOnlyCell(ws, value=v) for v in ("bucket", "region",
                                                    "gigs")])
    for (bucket, region), n in sorted(counts.items(), key=lambda kv: -kv[1]):
        ws.append([bucket, region, n])
    ws.append(["TOTAL", "", nrows])
    ws.append(["", "", ""])
    ws.append([WriteOnlyCell(ws, value=v) for v in ("source", "", "gigs")])
    for s, n in sorted(src_counts.items(), key=lambda kv: -kv[1]):
        ws.append([s, "", n])

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 112
    lines = [
        ("Personal pipeline — contract & temp", f"Snapshot {SNAP} (UTC). "
         f"{nrows} gigs: every Remote contract/temp/freelance role across all "
         "harvested sources, plus every on-site/hybrid gig in MA, RI, and the "
         "Miami / Fort Lauderdale / West Palm Beach metros."),
        ("flag", f"Pick {FLAG_CHOICES.replace(',', ' / ')} from the dropdown "
         "(free text works too). Interested rows turn green. Filter the flag "
         "column to see only your picks."),
        ("notes", "Free text — next steps, contacts, status."),
        ("bucket", "Remote (work from anywhere) or Local (on-site in your "
         "regions)."),
        ("region", "MA / RI / South FL. 'FL (statewide tag)' = posting said "
         "only 'Florida' with no city. Remote rows show a region when the "
         "employer also tagged a local office."),
        ("source / employment_type", "Which feed and what contingent type — "
         "filter to, say, only Contract, or only staffing_ats sources."),
        ("caveat", "Locations are freeform across dozens of feeds; region "
         "matching is best-effort (unambiguous cities match bare, ambiguous "
         "ones need an FL tag). Non-US country codes never match a local "
         "region."),
        ("everything else", "Columns carried from contract_master.csv — see "
         "the README in Contract_Gigs for the full glossary and the Sources "
         "tab for every source and how to reach it."),
        ("refresh", "Re-run the fetchers + contract_build.py + "
         "contract_pipeline.py. Re-apply flags to a new snapshot by matching "
         "apply_url or (company,title)."),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)

    # ---- splice the flag dropdown into the Pipeline sheet XML
    dv_xml = (f'<dataValidations count="1">'
              f'<dataValidation type="list" allowBlank="1" '
              f'showErrorMessage="0" sqref="A2:A{nrows + 1}">'
              f'<formula1>"{FLAG_CHOICES}"</formula1>'
              f'</dataValidation></dataValidations>')
    tmp = OUT_XLSX + ".tmp"
    with zipfile.ZipFile(OUT_XLSX) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                anchor = xml.rfind("</conditionalFormatting>")
                if anchor == -1:
                    raise SystemExit("conditionalFormatting anchor missing")
                pos = anchor + len("</conditionalFormatting>")
                xml = xml[:pos] + dv_xml + xml[pos:]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, OUT_XLSX)

    with zipfile.ZipFile(OUT_XLSX) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    i_dv = xml.find("<dataValidations")
    i_cf = xml.rfind("</conditionalFormatting>")
    i_pm = xml.find("<pageMargins")
    ok = (i_dv != -1 and i_cf != -1 and i_dv > i_cf
          and (i_pm == -1 or i_dv < i_pm)
          and xml.count("<dataValidations") == 1)
    print(f"verify: dropdown present and schema-ordered: {ok}", flush=True)
    if not ok:
        raise SystemExit("dataValidations splice verification failed")


if __name__ == "__main__":
    main()
