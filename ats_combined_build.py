#!/usr/bin/env python3
"""
ats_combined_build.py — Getro + Lever + Greenhouse + Ashby ONLY, one workbook.

Merges the four ATS/ecosystem master CSVs (no contract, portal, Bullhorn,
Phenom or job-aggregator sources) using the same canonical-ATS dedupe as
all_jobs_build.py, and writes ATS_Jobs_<date>.xlsx:

  README     what this is and how to read it
  All Jobs   every ecosystem job, newest first ("All Jobs 2" spill tab only
             if needed to stay under Excel's 1,048,576-row cap)
  PDL Jobs   jobs seen ONLY on Getro's People Data Labs firehose board,
             split out (as in the monolith) so the real-board tabs stay
             browsable and under the row cap
  Summary    counts by source / work mode / employment type / seniority
  Boards     jobs per board, newest post date

Same posted-date cutoff convention as the monolith: jobs first seen before
2026-05-01 are cut (--since YYYY-MM-DD to move it, --all-dates for no cut).

Columns that only the portal crawl ever fills (city, state) are dropped.

  python3 ats_combined_build.py                 last-cutoff default
  python3 ats_combined_build.py --all-dates     no posted-date cut
  python3 ats_combined_build.py --since 2026-06-01
"""

import csv
import os
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

import all_jobs_build as ab

BASE = os.path.dirname(os.path.abspath(__file__))
SNAP = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"ATS_Jobs_{SNAP}.xlsx")

POSTED_SINCE = "" if "--all-dates" in sys.argv \
    else ab._argval("--since", "2026-05-01")

# (label, master csv, loader, README blurb) — order drives load + README order.
SOURCES = [
    ("getro", ab.GETRO_CSV, ab.load_getro,
     "Getro VC/network ecosystem boards (~768 live)"),
    ("lever", ab.LEVER_CSV, ab.load_lever,
     "public Lever board crawl (~2,842 boards)"),
    ("greenhouse", ab.GREENHOUSE_CSV, ab.load_greenhouse,
     "public Greenhouse board crawl (~8,376 boards)"),
    ("ashby", ab.ASHBY_CSV, ab.load_ashby,
     "public Ashby board crawl (~4k boards)"),
]

DROP = {"city", "state"}
KEEP = [i for i, c in enumerate(ab.COLUMNS) if c not in DROP]
OUT_COLS = [ab.COLUMNS[i] for i in KEEP]
OUT_WIDTHS = [ab.WIDTHS[i] for i in KEEP]

HDR_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)


def hdr(ws, names):
    out = []
    for nm in names:
        c = WriteOnlyCell(ws, value=nm)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        out.append(c)
    return out


def out_row(r):
    return [ab.safe(r[i]) for i in KEEP]


def data_tabs(wb, base_name, rows):
    """Stream `rows` (raw row objects) across '<name>', '<name> 2', ...
    Each row is transformed with out_row() one at a time and handed straight
    to the write-only sheet, so the full ~1M-row transformed table is never
    materialized at once — that doubling is what OOM-kills a 16GB box."""
    names = []
    total = len(rows)
    chunks = max(1, -(-total // ab.XLSX_CAP)) if total else 1
    for ci in range(chunks):
        name = base_name if ci == 0 else f"{base_name} {ci + 1}"
        ws = wb.create_sheet(name)
        for j, wd in enumerate(OUT_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(j)].width = wd
        ws.freeze_panes = "C2"
        lo, hi = ci * ab.XLSX_CAP, (ci + 1) * ab.XLSX_CAP
        n_chunk = min(hi, total) - lo
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(OUT_COLS))}"
                              f"{n_chunk + 1}")
        ws.append(hdr(ws, OUT_COLS))
        for r in rows[lo:hi]:
            ws.append(out_row(r))
        names.append(name)
    return names


def main():
    csv.field_size_limit(10 ** 7)

    missing = [os.path.basename(p) for _, p, _, _ in SOURCES
               if not os.path.exists(p)]
    if missing:
        raise SystemExit(f"missing master CSV(s): {', '.join(missing)}")

    by_key = {}
    raw_per = {}
    merges = 0
    for pname, path, loader, _blurb in SOURCES:
        n = 0
        for r in loader(path):
            n += 1
            if n % 250_000 == 0:
                print(f"  ...{pname} {n} rows", flush=True)
            k, _tier = ab.job_key(r)
            prev = by_key.get(k)
            if prev is not None:
                ab.merge(prev, r)
                merges += 1
            else:
                r[ab.IDX["job_key"]] = k
                by_key[k] = r
        raw_per[pname] = n
        print(f"loaded {pname}: {n} rows (unified so far {len(by_key)})",
              flush=True)

    rows = list(by_key.values())
    by_key = None
    for r in rows:
        names = sorted(r[ab.BRDS])
        r[ab.IDX["board_count"]] = max(r[ab.BASECNT], len(names))
        r[ab.IDX["boards"]] = "; ".join(names[:6]) + (
            f" (+{len(names) - 6} more)" if len(names) > 6 else "")
        srcs = sorted(r[ab.SRCS])
        r[ab.IDX["sources"]] = ", ".join(srcs)
        r[ab.IDX["n_sources"]] = len(srcs)

    n_cut = 0
    if POSTED_SINCE:
        pre = len(rows)
        rows = [r for r in rows
                if not r[ab.IDX["posted"]]
                or r[ab.IDX["posted"]] >= POSTED_SINCE]
        n_cut = pre - len(rows)
        print(f"posted-date cutoff {POSTED_SINCE}: cut {n_cut}, "
              f"kept {len(rows)}", flush=True)
    rows.sort(key=lambda r: (r[ab.IDX["posted"]] or "",
                             r[ab.IDX["n_sources"]]), reverse=True)
    multi = sum(1 for r in rows if len(r[ab.SRCS]) > 1)
    total_raw = sum(raw_per.values())
    print(f"raw {total_raw} -> unified {len(rows)} "
          f"({merges} duplicate rows merged; {multi} jobs span 2+ sources)",
          flush=True)

    split = len(rows) > ab.XLSX_CAP
    if split:
        eco, pdl = [], []
        for r in rows:
            (pdl if (r[ab.SRCS] == {"getro"}
                     and r[ab.BRDS] == {ab.PDL_BOARD})
             else eco).append(r)
        print(f"split: {len(eco)} ecosystem rows, {len(pdl)} PDL-only rows",
              flush=True)
    else:
        eco, pdl = rows, []
        print(f"no split needed: {len(rows)} rows fit one tab", flush=True)

    by_source, by_wm, by_type, by_sen = {}, {}, {}, {}
    n_remote = 0
    board_stats = {}
    for r in rows:
        s = r[ab.IDX["sources"]] or "(none)"
        by_source[s] = by_source.get(s, 0) + 1
        wm = r[ab.IDX["work_mode"]] or "(not stated)"
        by_wm[wm] = by_wm.get(wm, 0) + 1
        et = r[ab.IDX["employment_type"]] or "(not stated)"
        by_type[et] = by_type.get(et, 0) + 1
        sen = r[ab.IDX["seniority"]] or "(not stated)"
        by_sen[sen] = by_sen.get(sen, 0) + 1
        if r[ab.IDX["remote"]] == "True":
            n_remote += 1
        for bn in r[ab.BRDS]:
            st = board_stats.setdefault(bn, [0, ""])
            st[0] += 1
            if (r[ab.IDX["posted"]] or "") > st[1]:
                st[1] = r[ab.IDX["posted"]] or ""

    wb = Workbook(write_only=True)

    lines = [
        ("GETRO + LEVER + GREENHOUSE + ASHBY — combined jobs, nothing else",
         ""),
        ("Snapshot (UTC)", SNAP),
        ("", ""),
        ("Unique jobs (deduped)", len(rows)),
    ]
    for pname, _path, _loader, _blurb in SOURCES:
        lines.append((f"  from {pname.capitalize()} master",
                      raw_per.get(pname, 0)))
    lines += [
        ("  duplicate rows merged away", merges),
        ("  jobs seen on 2+ of the four sources", multi),
        ("Remote jobs", n_remote),
        ("", ""),
        ("What this is", "Every job from four board universes merged into "
         "one deduped table: "
         + "; ".join(b for _, _, _, b in SOURCES)
         + ". The contract/gig crawl, the FT/contract portals (Bullhorn, "
           "Phenom, BambooHR, SmartRecruiters, Recruitee, Breezy) and the "
           "remote-job aggregators are NOT included in this workbook."),
        ("All Jobs tab", "Every job, newest first." if not split else
         "Excel caps a tab at 1,048,576 rows, so 'All Jobs' holds every job "
         "EXCEPT those seen only on Getro's People Data Labs firehose board "
         "(a LinkedIn-scale aggregator), which sit in 'PDL Jobs'. A job on "
         "PDL *and* any real board counts as ecosystem and stays in "
         "All Jobs. Both tabs are sorted newest first."),
        ("How deduping works", "The same posting is one row even when it "
         "appears on many boards or in several crawls. Match order: "
         "(1) canonical ATS identity parsed from the apply URL (platform + "
         "org + job id); (2) company + title + apply URL; (3) company + "
         "title + location. 'sources' lists which crawl(s) carried it; "
         "'boards' / 'board_count' list where it was posted. A Greenhouse or "
         "Ashby company that also sits on a Getro VC board collapses to one "
         "row carrying both sources."),
        ("Posted-date cutoff",
         (f"Jobs first seen before {POSTED_SINCE} are excluded "
          f"({n_cut} cut this snapshot). Rebuild with --since YYYY-MM-DD "
          "to move the line, or --all-dates for everything.") if POSTED_SINCE
         else "None — every job ever seen is included (--all-dates)."),
        ("job_key", "Stable dedupe key — match on it to carry notes into "
         "the next snapshot."),
        ("Refresh", "python3 ats_combined_build.py  (rebuilds this workbook "
         "from the four master CSVs already on disk)."),
    ]
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 116
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = BOLD
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    data_tabs(wb, "All Jobs", eco)
    if split:
        data_tabs(wb, "PDL Jobs", pdl)

    ws = wb.create_sheet("Summary")
    for col, wd in zip("AB", (34, 14)):
        ws.column_dimensions[col].width = wd

    def section(title, d):
        ws.append(hdr(ws, [title, "jobs"]))
        for k, v in sorted(d.items(), key=lambda kv: -kv[1]):
            ws.append([k, v])
        ws.append([])

    section("by source", by_source)
    section("by work_mode", by_wm)
    section("by employment_type", by_type)
    section("by seniority", by_sen)
    ws.append([WriteOnlyCell(ws, value="TOTAL unique"), len(rows)])

    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABC", (52, 10, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    top = sorted(board_stats.items(), key=lambda kv: -kv[1][0])[:20000]
    ws.auto_filter.ref = f"A1:C{len(top) + 1}"
    ws.append(hdr(ws, ["board", "jobs", "newest_post"]))
    for bn, (n_, newest) in top:
        ws.append([ab.safe(bn), n_, newest])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
