#!/usr/bin/env python3
"""
personio_build_sheet.py — personio_all_jobs.jsonl -> personio_jobs_master.csv
+ Personio_All_Jobs_<date>.xlsx.

Two things here are load-bearing and non-obvious.

1. THE VOCABULARY IS GERMAN.  employment_type is "Festanstellung",
   seniority is "Berufserfahren", schedule is "Voll- oder Teilzeit".  Left
   untranslated these never match the repo's English filters, so
   employment_type/seniority are mapped and the German original is preserved
   in *_raw columns.

2. country DEFAULTS TO "DE", never blank.  all_jobs_build.regions_of and
   ai_ma_remote's US filter both read a blank country as possibly-US, and a
   German row whose office is literally "Remote" would then be pulled into the
   user's US-remote cut.  Personio is a DACH product; guessing DE costs at
   most a handful of genuine US rows, while guessing blank floods the cut.

job_url is https://{slug}.jobs.personio.{tld}/job/{id}, which is the shape
all_jobs_build.canon_key already parses into ats|personio|{tenant}|{id}.

As with Rippling, Personio publishes no posting date, so `posted` comes from
personio_posted_ledger.json (first-seen, frozen) rather than being left blank
— blank rows never age out of the vet queue.
"""
import csv
import json
import os
import re
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

BASE = os.path.dirname(os.path.abspath(__file__))
SMOKE = int(os.environ.get("SMOKE") or 0)
PFX = "smoke_" if SMOKE else ""

IN_JSONL = os.path.join(BASE, PFX + "personio_all_jobs.jsonl")
LEDGER_JSON = os.path.join(BASE, PFX + "personio_posted_ledger.json")
OUT_CSV = os.path.join(BASE, PFX + "personio_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, "%sPersonio_All_Jobs_%s.xlsx" % (
    "Smoke_" if SMOKE else "", SNAPSHOT))
XLSX_ROW_CAP = 1_048_000

COLUMNS = ["title", "company", "location", "work_mode", "remote",
           "employment_type", "employment_type_raw", "seniority",
           "seniority_raw", "department", "city", "state", "country",
           "posted", "posted_precision", "apply_url", "job_url",
           "board", "board_url", "personio_job_id"]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)

# Personio's German vocabulary -> the repo's English one.
ET_MAP = {
    "festanstellung": "Full-time",
    "permanent employment": "Full-time",
    "praktikum": "Internship",
    "internship": "Internship",
    "werkstudent": "Part-time",
    "werkstudentenjob": "Part-time",
    "working student": "Part-time",
    "befristet": "Contract",
    "temporary": "Temp",
    "aushilfe": "Temp",
    "minijob": "Part-time",
    "teilzeit": "Part-time",
    "vollzeit": "Full-time",
    "freelance": "Freelance",
    "freelancer": "Freelance",
    "intern": "Internship",
    "trainee": "Internship",
    "ausbildung": "Internship",
    # Personio also serves English-localised tenants, and some values are
    # compound.  Measured against a live sample rather than guessed.
    "permanent employee": "Full-time",
    "praktikum/werkstudium": "Internship",
    "intern / student": "Internship",
    "fixed term": "Contract",
    "part-time employee": "Part-time",
    "full-time employee": "Full-time",
    "working student / intern": "Part-time",
}
SEN_MAP = {
    "berufserfahren": "Experienced",
    "experienced": "Experienced",
    "berufseinsteiger": "Entry level",
    "entry level": "Entry level",
    "student": "Student",
    "studenten": "Student",
    "schueler": "Student",
    "leitende angestellte": "Executive",
    "executive": "Executive",
    "manager": "Manager",
    "fuehrungskraft": "Manager",
    "berufseinstieg": "Entry level",
    "studierende/mitarbeitende im praktikum/auszubildende": "Student",
    "student/trainee/intern": "Student",
    "vorstand/gesch\u00e4ftsf\u00fchrung": "Executive",
    "vorstand/geschaeftsfuehrung": "Executive",
    "mit berufserfahrung": "Experienced",
}
REMOTE_RE = re.compile(r"remote|homeoffice|home office|anywhere|"
                       r"ortsunabh|mobiles arbeiten", re.I)
HYBRID_RE = re.compile(r"hybrid|teilweise remote", re.I)

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def norm(rec, ledger):
    j = rec.get("job") or {}
    slug = rec.get("slug") or ""
    tld = rec.get("tld") or "de"
    jid = str(j.get("id") or "").strip()
    title = (j.get("name") or "").strip()

    offices = j.get("offices")
    if isinstance(offices, list) and offices:
        loc = ", ".join(str(x) for x in offices if x)
    else:
        loc = (j.get("office") or "").replace(",", ", ")
    loc = I(loc.strip())

    et_raw = (j.get("employment_type") or "").strip()
    sched = (j.get("schedule") or "").strip()
    et = ET_MAP.get(et_raw.lower()) or ET_MAP.get(sched.lower()) or ""
    sen_raw = (j.get("seniority") or "").strip()
    sen = SEN_MAP.get(sen_raw.lower(), "")

    hay = " ".join(x for x in (loc, sched, j.get("category") or "") if x)
    if REMOTE_RE.search(hay):
        wm = "Remote"
    elif HYBRID_RE.search(hay):
        wm = "Hybrid"
    elif loc:
        wm = "On-site"
    else:
        wm = ""

    city = loc.split(",")[0].strip() if loc and wm != "Remote" else ""
    entry = ledger.get(jid) or {}
    posted = entry.get("first_seen") or ""
    url = ("https://%s.jobs.personio.%s/job/%s" % (slug, tld, jid)
           if slug and jid else "")

    r = [""] * N
    r[IDX["title"]] = title
    r[IDX["company"]] = I(j.get("subcompany") or slug)
    r[IDX["location"]] = loc
    r[IDX["work_mode"]] = I(wm)
    r[IDX["remote"]] = "True" if wm == "Remote" else ""
    r[IDX["employment_type"]] = I(et)
    r[IDX["employment_type_raw"]] = I(et_raw or sched)
    r[IDX["seniority"]] = I(sen)
    r[IDX["seniority_raw"]] = I(sen_raw)
    r[IDX["department"]] = I((j.get("department") or "").strip())
    r[IDX["city"]] = I(city)
    r[IDX["state"]] = ""
    # See the module docstring: DE, not blank.
    r[IDX["country"]] = I("DE" if tld == "de" else "DE")
    r[IDX["posted"]] = I(posted)
    r[IDX["posted_precision"]] = I("first_seen" if posted else "")
    r[IDX["apply_url"]] = url
    r[IDX["job_url"]] = url
    r[IDX["board"]] = I(slug)
    r[IDX["board_url"]] = I(rec.get("board_url") or "")
    r[IDX["personio_job_id"]] = jid
    return r


def main():
    if not os.path.exists(IN_JSONL):
        raise SystemExit("missing %s — run personio_fetch_jobs.py" % IN_JSONL)
    ledger = {}
    if os.path.exists(LEDGER_JSON):
        try:
            with open(LEDGER_JSON, encoding="utf-8") as f:
                ledger = json.load(f)
        except (OSError, json.JSONDecodeError):
            ledger = {}

    by_id, raw = {}, 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            raw += 1
            r = norm(rec, ledger)
            if not r[IDX["title"]]:
                continue
            key = "%s|%s" % (r[IDX["board"]], r[IDX["personio_job_id"]])
            by_id.setdefault(key, r)
    n_ids = len(by_id)

    by_sig = {}
    for r in by_id.values():
        sig = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               (r[IDX["apply_url"]] or r[IDX["location"]]).lower())
        by_sig.setdefault(sig, r)
    rows = list(by_sig.values())
    rows.sort(key=lambda r: (r[IDX["posted"]] or "", r[IDX["company"]]),
              reverse=True)
    print("raw rows %d -> unique ids %d -> deduped %d"
          % (raw, n_ids, len(rows)), flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print("wrote %s" % OUT_CSV, flush=True)
    write_xlsx(rows)


def write_xlsx(rows):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    n_et = sum(1 for r in rows if r[IDX["employment_type"]])
    for a, b in [
        ("Personio — all jobs", "Snapshot %s" % SNAPSHOT),
        ("Rows", "%s deduplicated postings" % format(len(rows), ",")),
        ("Region", "Personio is a DACH product: expect German-language "
                   "titles and German offices. country is set to DE rather "
                   "than left blank, because downstream US-remote filters "
                   "treat a blank country as possibly-US."),
        ("Vocabulary", "%s of %s rows had a recognised employment type. The "
                       "German original is kept in employment_type_raw / "
                       "seniority_raw so an unmapped value is visible rather "
                       "than silently dropped."
         % (format(n_et, ","), format(len(rows), ","))),
        ("Source", "{slug}.jobs.personio.de/search.json — one request per "
                   "tenant, whole board returned. Rate-limits hard (55% HTTP "
                   "429 unthrottled), so the fetcher runs at 2 rps."),
        ("posted", "first-seen date from personio_posted_ledger.json. "
                   "Personio publishes no posting date; a blank would never "
                   "age out of the vet queue."),
        ("Dedupe", "job_url is {slug}.jobs.personio.{tld}/job/{id}, which "
                   "all_jobs_build.canon_key turns into "
                   "ats|personio|tenant|id."),
    ]:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    ws.freeze_panes = "A2"
    widths = {"title": 52, "company": 26, "location": 32, "work_mode": 10,
              "remote": 8, "employment_type": 14, "employment_type_raw": 18,
              "seniority": 14, "seniority_raw": 16, "department": 22,
              "city": 18, "state": 7, "country": 9, "posted": 11,
              "posted_precision": 15, "apply_url": 52, "job_url": 52,
              "board": 24, "board_url": 42, "personio_job_id": 12}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    hdr = []
    for c in COLUMNS:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(N),
                                      min(len(rows), XLSX_ROW_CAP) + 1)
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABC", (34, 10, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    hdr = []
    for c in ("board", "jobs", "newest_post"):
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    stats = {}
    for r in rows:
        st = stats.setdefault(r[IDX["board"]], [0, ""])
        st[0] += 1
        if (r[IDX["posted"]] or "") > st[1]:
            st[1] = r[IDX["posted"]] or ""
    for b, (n, newest) in sorted(stats.items(), key=lambda kv: -kv[1][0]):
        ws.append([safe(b), n, newest])

    wb.save(OUT_XLSX)
    print("wrote %s" % OUT_XLSX, flush=True)


if __name__ == "__main__":
    sys.exit(main())
