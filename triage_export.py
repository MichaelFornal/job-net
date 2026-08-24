#!/usr/bin/env python3
"""triage_export.py — what you approved, and how far you have got with it.

    python3 triage_export.py                 # Approved_<date>.xlsx + .csv
    python3 triage_export.py --maybe         # add a 'Maybe' tab
    python3 triage_export.py --harvest F.xlsx # read status edits back out of F

The workbook carries the full 36-column record, not the trimmed card that
triage.db stores, so it goes back to MA_US_Remote_<date>.csv for the row bodies
and uses triage.db only for the verdicts.

Two tabs, and the order they are created in matters (see `workbook`):

  * **Worklist** — the compact thing you actually work: status, company, title,
    where, when, pay, a live apply link, a note.  Sheet 1 because the status
    dropdown is spliced into `sheet1.xml` by file name.
  * **Approved** — the full 36-column record, unchanged from before.

The workbook is a *report*.  http://127.0.0.1:8766/apply owns the status: it
writes to triage.db the moment you click, and every workbook is regenerated
from that.  `--harvest` is the one path back in, for when you would rather tick
boxes in Excel — and it only ever **fills in**, never overwrites.  A job whose
status was set on the page is left exactly as it is, so regenerating can't lose
a click and harvesting can't lose one either.

"""
import csv
import datetime as dt
import glob
import os
import sys
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
sys.path.insert(0, os.path.join(BASE, "apply"))

import all_jobs_build as ajb                          # noqa: E402
import ma_remote_build as mrb                         # noqa: E402
import triage_load as TL                              # noqa: E402
from triage import apply_queue as AQ                  # noqa: E402
from triage import lanes as L                         # noqa: E402

SNAP = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d")
EXTRA = ["status", "applied_at", "note", "verdict", "decided_at", "lane"]
EXTRA_W = [11, 12, 30, 10, 20, 22]

# Excel shows the human label; triage.db stores the machine value.  Harvest
# accepts either, case-insensitively, so a hand-typed "applied" still lands.
LABEL = {"to_apply": "To apply", "applying": "Applying", "applied": "Applied",
         "skipped": "Skipped", "dead": "Dead"}
UNLABEL = {v.casefold(): k for k, v in LABEL.items()}
UNLABEL.update({k: k for k in LABEL})
UNLABEL[""] = "to_apply"

# The compact tab.  `job_key` is last and looks like noise, but it is the only
# thing --harvest can key a row on after you have sorted and filtered the sheet.
# `send` is which resume this row wants, resolved from the lane it was decided in
# through lanes.LANE_RESUME — the same constant the triage card prints, so the sheet
# and the UI can never recommend different documents for the same job. Blank for the
# unscored lanes (ai12, ma, rest), which have no swath resume of their own.
WORK_COLS = ["status", "send", "company", "title", "location", "posted", "salary",
             "source", "ai_band", "applied_at", "note", "apply_url", "job_key"]
WORK_W = [11, 22, 26, 46, 24, 11, 15, 13, 26, 12, 30, 62, 34]

HDR_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
# Same palette as vet_build.py, so a status means the same colour everywhere.
STATUS_FILLS = [
    ('$A2="Applied"',  "BDD7EE"),      # blue   — done
    ('$A2="Applying"', "FFE699"),      # amber  — in flight, not sent
    ('$A2="To apply"', "C6EFCE"),      # green  — live work
    ('OR($A2="Skipped",$A2="Dead")', "D9D9D9"),   # grey — closed out
]


def newest_csv():
    hits = sorted(glob.glob(os.path.join(BASE, "MA_US_Remote_*.csv")))
    if not hits:
        sys.exit("no MA_US_Remote_*.csv found")
    return hits[-1]


def verdicts(con, want):
    """{job_key: (status, applied_at, note, verdict, decided_at, lane)}.

    LEFT JOIN, because `applications` only holds jobs you have acted on —
    everything else is 'to apply' by absence (see triage/apply_queue.py).
    """
    q = ",".join("?" * len(want))
    return {r["job_key"]: (LABEL.get(r["status"], r["status"]),
                           r["applied_at"], r["note"],
                           r["verdict"], r["decided_at"], r["lane"])
            for r in con.execute(
                f"SELECT d.job_key, d.verdict, d.decided_at, d.lane, "
                f"       COALESCE(a.status, 'to_apply') AS status, "
                f"       COALESCE(a.applied_at, '')     AS applied_at, "
                f"       COALESCE(a.note, '')           AS note "
                f"  FROM decisions d "
                f"  LEFT JOIN applications a ON a.job_key = d.job_key "
                f" WHERE d.verdict IN ({q}) AND d.reason = 'manual'", want)}


def collect(paths, keys):
    """Full CSV rows for `keys`, newest first, grouped by verdict.

    Takes EVERY feed's CSV, not just the main corpus.  triage.db holds one
    `jobs` table across feeds, but the row bodies live in whichever file that
    feed was loaded from, so a single-file scan silently dropped approved
    contract gigs: approved in the app, absent from the workbook, with nothing
    anywhere saying a row had gone missing.

    The first path's header is canonical.  Later files are mapped onto it by
    column name and anything they don't have is written blank — a contract row
    has no startup_tier, and inventing one would put a value in a column the
    gig was never judged on.
    """
    csv.field_size_limit(10 ** 9)
    if isinstance(paths, str):
        paths = [paths]
    vi = EXTRA.index("verdict")      # not 0 — EXTRA leads with the apply status
    out = {}
    hdr = None
    for path in paths:
        if not path or not os.path.exists(path):
            continue
        with open(path, newline="", encoding="utf-8") as f:
            r = csv.reader(f)
            own = next(r)
            if hdr is None:
                hdr, pick = own, list(range(len(own)))
            else:
                idx = {c: i for i, c in enumerate(own)}
                pick = [idx.get(c) for c in hdr]
            ki, pi = own.index("job_key"), own.index("posted")
            for row in r:
                if len(row) < len(own):
                    continue
                v = keys.get(row[ki])
                if v:
                    out.setdefault(v[vi], []).append(
                        list(v) + [row[i] if i is not None else ""
                                   for i in pick])
    if hdr is None:
        sys.exit("no corpus CSV could be read")
    pi = hdr.index("posted")
    for v in out.values():
        v.sort(key=lambda x: (x[len(EXTRA) + pi] or ""), reverse=True)
    return out, hdr


def money(get):
    """'165k–210k' from the salary columns, or '' — the Worklist's pay cell."""
    def one(v):
        try:
            n = float(v)
        except (TypeError, ValueError):
            return None
        if n <= 0:
            return None
        return f"${round(n / 1000)}k" if n >= 1000 else f"${round(n)}"
    lo, hi = one(get("salary_min")), one(get("salary_max"))
    if not lo and not hi:
        return ""
    rng = lo if lo == hi or not hi else (hi if not lo else f"{lo}–{hi}")
    per = (get("salary_period") or "").lower()
    return rng + (f" /{per}" if per and per not in ("year", "yearly") else "")


def worklist(rows, hdr):
    """The compact rows, unworked first so the work is at the top."""
    idx = {c: i + len(EXTRA) for i, c in enumerate(hdr)}
    idx.update({c: i for i, c in enumerate(EXTRA)})
    out = []
    for row in rows:
        get = lambda c: row[idx[c]] if idx[c] < len(row) else ""  # noqa: E731
        res = L.resume_for((get("lane") or "").strip())
        out.append([get("status"), f"{res[0]} · {res[1]}" if res else "",
                    get("company"), get("title"),
                    get("location"), get("posted"), money(get),
                    get("source"), get("ai_band"), get("applied_at"),
                    get("note"), get("apply_url") or get("job_url"),
                    get("job_key")])
    # Same story the /apply page tells, in the same order: in-flight first,
    # then unworked, strongest AI signal first, newest first.  Two stable
    # passes rather than one tuple key because the keys sort in opposite
    # directions.  Anything you have half-finished goes above even the best
    # untouched row — it is the shortest path to one more application sent.
    order = {LABEL["applying"]: 0, LABEL["to_apply"]: 1}

    def rank(r):
        band = (r[8] or "").strip()[:1]      # ai_band, "1 Claude / agentic…"
        return int(band) if band.isdigit() else 9
    out.sort(key=lambda r: r[5] or "", reverse=True)          # posted DESC
    out.sort(key=lambda r: (order.get(r[0], 2), rank(r)))
    return out


def hdr_cells(ws, names, widths, fill=True):
    for j, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    out = []
    for nm in names:
        c = WriteOnlyCell(ws, value=nm)
        c.font = HDR_FONT if fill else Font(bold=True)
        if fill:
            c.fill = HDR_FILL
        out.append(c)
    return out


def worklist_tab(wb, rows):
    """Sheet 1 — and it has to stay sheet 1; splice_dropdown() finds it by the
    file name openpyxl gives it, which in write-only mode is creation order."""
    ws = wb.create_sheet("Worklist")
    # Before the first append: write-only mode serialises sheet properties with
    # the opening tag, so anything set after row 1 is silently dropped.
    ws.freeze_panes = "B2"
    ws.append(hdr_cells(ws, WORK_COLS, WORK_W))
    last = get_column_letter(len(WORK_COLS))
    if not rows:
        return 0            # an empty range is not a legal sqref
    ws.auto_filter.ref = f"A1:{last}{len(rows) + 1}"
    url_i = WORK_COLS.index("apply_url")
    for row in rows:
        cells = []
        for i, v in enumerate(row):
            c = WriteOnlyCell(ws, value=ajb.safe(v))
            if i == url_i and v:
                c.hyperlink = v
                c.font = LINK_FONT
            cells.append(c)
        ws.append(cells)
    for formula, color in STATUS_FILLS:
        ws.conditional_formatting.add(
            f"A2:{last}{len(rows) + 1}",
            FormulaRule(formula=[formula],
                        fill=PatternFill("solid", start_color=color,
                                         end_color=color),
                        stopIfTrue=False))
    return len(rows)


def splice_dropdown(path, nrows):
    """Add the status dropdown to Worklist!A.

    openpyxl's write-only mode cannot emit DataValidation, so the sheet XML is
    edited in the saved zip — the same trick as vet_build.py.  Two schema
    facts this depends on: the sheet is `sheet1.xml` because it was created
    first, and <dataValidations> must sit after </conditionalFormatting> and
    before <pageMargins>, which is why it is anchored rather than appended.
    """
    target = "xl/worksheets/sheet1.xml"
    dv = (f'<dataValidations count="1"><dataValidation type="list" '
          f'allowBlank="1" showErrorMessage="0" sqref="A2:A{nrows + 1}">'
          f'<formula1>"{",".join(LABEL.values())}"</formula1>'
          f'</dataValidation></dataValidations>')
    tmp = path + ".tmp"
    seen = False
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == target:
                seen = True
                xml = data.decode("utf-8")
                anchor = xml.rfind("</conditionalFormatting>")
                if anchor == -1:
                    raise SystemExit(f"{target}: conditionalFormatting anchor "
                                     "missing — is Worklist still sheet 1?")
                pos = anchor + len("</conditionalFormatting>")
                data = (xml[:pos] + dv + xml[pos:]).encode("utf-8")
            zout.writestr(item, data)
    if not seen:
        os.remove(tmp)
        raise SystemExit(f"{target} not in the workbook")
    os.replace(tmp, path)

    with zipfile.ZipFile(path) as z:                  # self-verify
        xml = z.read(target).decode("utf-8")
    if xml.count("<dataValidations") != 1:
        raise SystemExit("dropdown spliced more than once")
    if not (xml.index("</conditionalFormatting>") < xml.index("<dataValidations")
            < xml.index("<pageMargins")):
        raise SystemExit("dataValidations landed in the wrong place")


def workbook(groups, hdr, out_xlsx):
    cols = EXTRA + hdr
    widths = EXTRA_W + mrb.WIDTHS
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                        fill_type="solid")

    n_work = worklist_tab(wb, worklist(groups.get("approve", []), hdr))

    for name in ("approve", "maybe"):
        rows = groups.get(name)
        if not rows:
            continue
        for n, chunk in enumerate(
                [rows[i:i + ajb.XLSX_CAP]
                 for i in range(0, len(rows), ajb.XLSX_CAP)] or [[]], 1):
            title = ("Approved" if name == "approve" else "Maybe")
            ws = wb.create_sheet(title if n == 1 else f"{title} {n}")
            for j, wd in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(j)].width = wd
            ws.freeze_panes = "G2"
            ws.auto_filter.ref = \
                f"A1:{get_column_letter(len(cols))}{len(chunk) + 1}"
            head = ajb.hdr(ws, cols, bold)
            if name == "approve":
                for c in head[:len(EXTRA)]:
                    c.fill = green
            ws.append(head)
            for row in chunk:
                ws.append([ajb.safe(v) for v in row])
    wb.save(out_xlsx)
    if n_work:
        # No rows means no conditional formatting, which means no anchor to
        # splice against — and an "is Worklist still sheet 1?" error that sends
        # you hunting the wrong invariant.
        splice_dropdown(out_xlsx, n_work)


# ------------------------------------------------------------------- harvest
def harvest(con, path):
    """Read status edits back out of a workbook's Worklist tab.

    **Fills in; never overwrites.**  A job carrying a status in `applications`
    is left alone, so this can be run against a stale workbook without
    reverting anything you clicked since.  No timestamp comparison anywhere.

    The one gap, stated plainly because it cannot be closed without giving up
    "'to apply' is the absence of a row": a job you deliberately reset to
    To apply on the page has no row, and so is indistinguishable from one you
    never touched.  Harvesting an older workbook that still says "Applied"
    will therefore re-apply it.  Harvest first, click second, and it can't
    bite you.

    Columns are found by header name, not position, so sorting and filtering
    the sheet in Excel before saving is safe.
    """
    if not os.path.exists(path):
        sys.exit(f"no such workbook: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Worklist" not in wb.sheetnames:
        sys.exit(f"{os.path.basename(path)} has no 'Worklist' tab — "
                 f"tabs are: {', '.join(wb.sheetnames)}")
    rows = wb["Worklist"].iter_rows(values_only=True)
    norm = lambda v: ("" if v is None else str(v)).strip()   # noqa: E731
    head = {norm(v).casefold(): i for i, v in enumerate(next(rows, ()) or ())}
    for need in ("job_key", "status"):
        if need not in head:
            sys.exit(f"Worklist is missing a '{need}' column")
    have_note = "note" in head

    cur = {r["job_key"]: r["status"] for r in con.execute(
        "SELECT job_key, status FROM applications")}
    # A workbook can outlive a corpus rebuild, and its job_keys go dead when
    # one moves the job to a new key.  Writing to a dead key "succeeds" and is
    # then invisible to every live-filtered read — and worse, it gives the dead
    # row a newer updated_at than the live one, which is enough to make
    # carry_forward_applications prefer the stale status at the next rebuild.
    # So resolve each key to its live sibling by ct_key before writing.
    live = {r["job_key"] for r in con.execute(
        "SELECT job_key FROM jobs WHERE live = 1")}
    by_ct = {r["ct_key"]: r["job_key"] for r in con.execute(
        "SELECT ct_key, job_key FROM jobs WHERE live = 1 AND ct_key <> ''")}
    ct_of = {r["job_key"]: r["ct_key"] for r in con.execute(
        "SELECT job_key, ct_key FROM jobs")}

    took, held, blank, unknown, stale, moved = 0, 0, 0, 0, 0, 0
    for vals in rows:
        get = lambda c: norm(vals[head[c]]) if head[c] < len(vals) else ""  # noqa: E731
        key = get("job_key")
        if not key:
            continue
        want = UNLABEL.get(get("status").casefold())
        if want is None:
            unknown += 1
            continue
        if key not in live:                  # the workbook predates a rebuild
            alt = by_ct.get(ct_of.get(key, "\0"))
            if not alt:
                stale += 1
                continue
            key = alt
            moved += 1
        if key in cur:                       # the page already owns this row
            held += 1
            continue
        if want == "to_apply":
            blank += 1
            continue
        # Count what the database actually changed, never the request.
        took += 1 if AQ.set_status(
            con, key, want, get("note") if have_note else None) else 0

    print(f"harvested {os.path.basename(path)}: {took:,} status "
          f"{'change' if took == 1 else 'changes'} taken in")
    if moved:
        print(f"  {moved:,} followed the job to a new job_key since export")
    if held:
        print(f"  {held:,} left alone — already set on the /apply page")
    if blank:
        print(f"  {blank:,} still 'To apply' in the sheet")
    if stale:
        print(f"  {stale:,} are no longer in the corpus at all — skipped")
    if unknown:
        print(f"  {unknown:,} rows had a status that isn't one of "
              f"{', '.join(LABEL.values())}")
    if took:
        print("\nre-run `python3 triage_export.py` to regenerate the workbook "
              "with these folded in")
    return took


def main():
    args = sys.argv[1:]
    if "-h" in args or "--help" in args:
        print(__doc__)
        return
    want = ["approve"] + (["maybe"] if "--maybe" in args else [])
    # An explicit --csv means "export exactly this"; otherwise take every feed
    # the app can queue, so nothing you approved is missing from the workbook.
    if "--csv" in args:
        src = [args[args.index("--csv") + 1]]
    else:
        src = [newest_csv()] + sorted(
            glob.glob(os.path.join(BASE, "Contract_Pool_*.csv")))[-1:]

    con = TL.init()          # init, not connect: guarantees `applications`

    if "--harvest" in args:
        # Harvest and stop, deliberately: regenerating in the same run would
        # overwrite the workbook you just edited, and if anything about the
        # harvest looked wrong you would have nothing left to compare against.
        i = args.index("--harvest") + 1
        if i >= len(args):
            sys.exit("--harvest needs a workbook path")
        harvest(con, args[i])
        con.close()
        return

    keys = verdicts(con, want)
    if not keys:
        sys.exit("nothing approved yet — run: python3 triage/server.py")
    groups, hdr = collect(src, keys)

    approved = groups.get("approve", [])
    out_csv = os.path.join(BASE, f"Approved_{SNAP}.csv")
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(EXTRA + hdr)
        w.writerows(approved)

    out_xlsx = os.path.join(BASE, f"Approved_{SNAP}.xlsx")
    workbook(groups, hdr, out_xlsx)

    ac = AQ.counts(con)
    print(f"{len(approved):,} approved"
          + (f" · {len(groups.get('maybe', [])):,} maybe" if "--maybe" in args
             else ""))
    print(f"  {ac['applied']:,} applied"
          + (f" · {ac['applying']:,} applying" if ac["applying"] else "")
          + f" · {ac['to_apply']:,} to apply"
          + (f" · {ac['skipped']:,} skipped" if ac["skipped"] else "")
          + (f" · {ac['dead']:,} dead" if ac["dead"] else ""))
    missing = len(keys) - sum(len(v) for v in groups.values())
    if missing:
        print(f"  note: {missing:,} decided job_keys are not in "
              f"{os.path.basename(src)} — they aged out of the corpus. Their "
              f"verdicts are still in triage.db.")
    print(f"wrote {out_csv}")
    print(f"wrote {out_xlsx} ({os.path.getsize(out_xlsx) / 1e6:.1f} MB)")

    con.close()


if __name__ == "__main__":
    main()
