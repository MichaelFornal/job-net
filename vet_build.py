#!/usr/bin/env python3
"""
vet_build.py — the vetting workbook: filter + vet Getro/Lever jobs, keep the
decisions across rebuilds.

One command:  python3 vet_build.py

  1. HARVEST   reads your newest Vet_Queue_*.xlsx and banks every row with vet
               data (status/notes/applied_date/contact/next_action/
               follow_up_date) into vet_ledger.csv, keyed on job_key.
  2. REBUILD   merges getro_jobs_master.csv + lever_jobs_master.csv +
               greenhouse_jobs_master.csv (same canonical-ATS dedupe as
               all_jobs_build.py; greenhouse skipped with a warning if its
               master is missing) and slices to the vettable surface:
               US-remote jobs plus on-site/hybrid roles in MA / RI / S-FL.
  3. MERGE     re-applies the ledger onto the new snapshot. Vetted jobs never
               age out of the posted-date cutoff; jobs with an active status
               (anything but No/Rejected) that vanished from the crawl are
               kept as 'gone' rows so an application can't silently disappear.
  4. WRITE     Vet_Queue_<date>.xlsx — Queue (status dropdown, color-coded),
               Top Roles (the best ~250 AI-builder roles in the slice, ranked
               best-first), README, Summary — then records this queue's
               job_keys in vet_seen_keys.txt so next build can star what's NEW.

The Top Roles tab ranks by AI-builder strength: Claude/Anthropic/agentic and
forward-deployed titles first, then core AI engineer/LLM/GenAI, then AI-ML
adjacent, then engineering titles whose skills list AI tooling — newest first
within each band. Vet on either tab; status carries to both (both are harvested).

Flags: --since YYYY-MM-DD (posted cutoff, default 2026-05-01), --all-dates,
       --top N (how many Top Roles to rank, default 250),
       --harvest-only (bank vet work, skip the rebuild),
       --harvest-from PATH (harvest a specific workbook instead of newest).
"""

import csv
import datetime as dt
import gc
import glob
import os
import re
import sys
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

import all_jobs_build as ab

BASE = os.path.dirname(os.path.abspath(__file__))
SNAP = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Vet_Queue_{SNAP}.xlsx")
LEDGER = os.path.join(BASE, "vet_ledger.csv")
SEEN_KEYS = os.path.join(BASE, "vet_seen_keys.txt")

POSTED_SINCE = "" if "--all-dates" in sys.argv \
    else ab._argval("--since", "2026-05-01")

VET_COLS = ["status", "notes", "applied_date", "contact", "next_action",
            "follow_up_date"]
CTX_COLS = ["new", "listing", "bucket", "region"]
DROP = {"city", "state"}
KEEP = [i for i, c in enumerate(ab.COLUMNS) if c not in DROP]
OUT_COLS = [ab.COLUMNS[i] for i in KEEP]
QCOLS = VET_COLS + CTX_COLS + OUT_COLS
QWIDTHS = [12, 28, 12, 18, 18, 13, 5, 8, 9, 16] + [ab.WIDTHS[i] for i in KEEP]

# Top Roles tab: same vet columns, plus rank / ai_band / matched context.
TOP_CTX = ["rank", "ai_band", "matched", "new", "bucket", "region"]
TOPCOLS = VET_COLS + TOP_CTX + OUT_COLS
TOPWIDTHS = [12, 28, 12, 18, 18, 13, 5, 17, 22, 5, 9, 16] \
    + [ab.WIDTHS[i] for i in KEEP]
TOP_N = int(ab._argval("--top", "250"))

STATUSES = ["Interested", "Maybe", "No", "Applied", "Interviewing", "Offer",
            "Rejected"]
INACTIVE = {"No", "Rejected"}
US_OK = ("US", "USA", "UNITED STATES", "")

LEDGER_COLS = ["job_key"] + VET_COLS + ["company", "title", "apply_url",
                                        "posted", "first_seen", "last_updated"]
BUCKET_RANK = {"Remote": 0, "Local": 1, "Other": 2, "Gone": 3}

# ---- AI-builder banding for the Top Roles tab (best-first). Higher band =
# stronger "this job is about building with AI" signal. Regexes adapted from
# ai_jobs_filter.py; the Claude/Anthropic/agentic/forward-deployed hits are
# split into their own top band since those are the bullseye.
AI_CLAUDE = re.compile(r"""(?ix)
    \bclaude\b | \banthropic\b | \bagentic\b | \bai\s+agents?\b |
    \bforward[\s-]deployed | \bcoding\s+agents?\b | \bai[\s-]native\b |
    \bfounding\s+ai\b""")
AI_CORE = re.compile(r"""(?ix)
    \bllms?\b | \blarge\s+language\s+model | \bgen\s?ai\b |
    \bgenerative\s+ai\b | \bapplied\s+ai\b | \bprompt\s+engineer |
    \bai[\s/-]+(?:engineer|developer|dev\b|builder|application|apps?\b|
        product\s+engineer|full[\s-]?stack|software|solutions?\s+engineer|
        automation|platform\s+engineer|integration|tooling) |
    \bopenai\b | \bgpt\b | \blangchain\b | \bRAG\b |
    \bcopilot\s+(?:engineer|developer)""")
AI_ADJ = re.compile(r"""(?ix)
    \bmachine\s+learning\b | \bml\s+(?:engineer|scientist|ops|platform|infra) |
    \bmlops\b | \bai/ml\b | \bml/ai\b | \bartificial\s+intelligence\b |
    \bdeep\s+learning\b | \bnlp\b | \bconversational\s+ai\b | \bchatbot |
    \bcomputer\s+vision\b |
    \bai\s+(?:architect|scientist|researcher|research|consultant|specialist|
        lead|manager|strategist)\b | \bcopilot\b | \bhugging\s?face\b""")
AI_SKILL = re.compile(r"""(?ix)
    \bllms?\b | \blangchain\b | \bllamaindex\b | \bopenai\b | \banthropic\b |
    \bclaude\b | \bgpt\b | \bretrieval[-\s]augmented | \bRAG\b |
    prompt\s+engineering | generative\s+ai | \bgen\s?ai\b | agentic |
    hugging\s?face | vector\s+database | ai\s+coding | copilot""")
AI_ENG_TITLE = re.compile(
    r"(?i)engineer|developer|\bswe\b|software|full[ -]?stack|back[ -]?end|"
    r"front[ -]?end|architect|programmer|\bdev\b")
AI_BANDS = {5: "Claude / Agentic", 4: "Core AI building",
            3: "AI / ML adjacent", 2: "AI skills match"}


def ai_band(title, skills):
    """(band, matched_keyword) for a role; band 0 = not an AI-builder role."""
    for band, rx in ((5, AI_CLAUDE), (4, AI_CORE), (3, AI_ADJ)):
        m = rx.search(title)
        if m:
            return band, re.sub(r"\s+", " ", m.group(0)).strip()
    if skills and AI_ENG_TITLE.search(title):
        m = AI_SKILL.search(skills)
        if m:
            return 2, re.sub(r"\s+", " ", m.group(0)).strip()
    return 0, ""


HDR_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = [
    ('$A2="Interested"', "C6EFCE"),
    ('$A2="Maybe"', "FFF2CC"),
    ('$A2="Applied"', "BDD7EE"),
    ('$A2="Interviewing"', "DDEBF7"),
    ('$A2="Offer"', "FFD966"),
    ('OR($A2="No",$A2="Rejected")', "D9D9D9"),
]
GONE_FILL = ('$H2="gone"', "F8CBAD")  # H = 'listing' column, Queue tab only


def hdr(ws, names):
    out = []
    for nm in names:
        c = WriteOnlyCell(ws, value=nm)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        out.append(c)
    return out


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


# ---- ledger -----------------------------------------------------------------
def load_ledger():
    led = {}
    if os.path.exists(LEDGER):
        with open(LEDGER, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                led[row["job_key"]] = {c: row.get(c, "") for c in LEDGER_COLS}
    return led


def save_ledger(led):
    tmp = LEDGER + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LEDGER_COLS)
        w.writeheader()
        for row in sorted(led.values(),
                          key=lambda r: (r["company"], r["title"])):
            w.writerow(row)
    os.replace(tmp, LEDGER)


def newest_queue():
    cands = [p for p in glob.glob(os.path.join(BASE, "Vet_Queue_*.xlsx"))
             if not os.path.basename(p).startswith("~$")]
    return max(cands, key=os.path.getmtime) if cands else ""


def harvest(led, path):
    print(f"harvesting {os.path.basename(path)} ...", flush=True)
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [s for s in ("Queue", "Top Roles") if s in wb.sheetnames]
    if not sheets:
        print("  no vettable sheet — skipped", flush=True)
        wb.close()
        return 0
    # A job can appear in BOTH tabs; merge so a blank appearance in one never
    # wipes vet data entered in the other. Non-empty wins; blank-in-all clears.
    merged = {}   # job_key -> {"vet": {...}, "ctx": (...), "has": bool}
    for sheet in sheets:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        head = {norm(v): i for i, v in enumerate(next(rows, ()) or ())}
        need = VET_COLS + ["job_key", "company", "title", "apply_url", "posted"]
        missing = [c for c in need if c not in head]
        if missing:
            raise SystemExit(f"harvest: {path} '{sheet}' lacks columns "
                             f"{missing}")
        for vals in rows:
            get = lambda c: norm(vals[head[c]]) if head[c] < len(vals) else ""
            jk = get("job_key")
            if not jk:
                continue
            vet = {c: get(c) for c in VET_COLS}
            has = any(vet.values())
            cur = merged.get(jk)
            if cur is None:
                merged[jk] = {"vet": vet, "has": has,
                              "ctx": (get("company"), get("title"),
                                      get("apply_url"), get("posted"))}
            elif has and not cur["has"]:
                cur["vet"], cur["has"] = vet, True
    n_vet = 0
    for jk, rec in merged.items():
        if not rec["has"] and jk not in led:
            continue          # untouched row, nothing to bank or clear
        prev = led.get(jk)
        company, title, apply_url, posted = rec["ctx"]
        row = {"job_key": jk, "company": company, "title": title,
               "apply_url": apply_url, "posted": posted,
               "first_seen": prev["first_seen"] if prev else SNAP,
               "last_updated": SNAP}
        row.update(rec["vet"])
        led[jk] = row
        n_vet += 1
    wb.close()
    print(f"  banked {n_vet} vetted rows (ledger now {len(led)})", flush=True)
    return n_vet


# ---- build ------------------------------------------------------------------
def load_rows():
    csv.field_size_limit(10 ** 7)
    by_key = {}
    for pname, path, loader in [
            ("getro", ab.GETRO_CSV, ab.load_getro),
            ("lever", ab.LEVER_CSV, ab.load_lever),
            ("greenhouse", ab.GREENHOUSE_CSV, ab.load_greenhouse)]:
        if not os.path.exists(path):
            if pname == "greenhouse":  # newest source — build without it
                print(f"WARNING: {os.path.basename(path)} missing — "
                      "greenhouse not included", flush=True)
                continue
            raise SystemExit(f"{os.path.basename(path)} missing")
        n = 0
        for r in loader(path):
            n += 1
            if n % 250_000 == 0:
                print(f"  ...{pname} {n} rows", flush=True)
            k, _tier = ab.job_key(r)
            prev = by_key.get(k)
            if prev is not None:
                ab.merge(prev, r)
            else:
                r[ab.IDX["job_key"]] = k
                by_key[k] = r
        print(f"loaded {pname}: {n} rows (unified so far {len(by_key)})",
              flush=True)
    rows = list(by_key.values())
    for r in rows:
        names = sorted(r[ab.BRDS])
        r[ab.IDX["board_count"]] = max(r[ab.BASECNT], len(names))
        r[ab.IDX["boards"]] = "; ".join(names[:6]) + (
            f" (+{len(names) - 6} more)" if len(names) > 6 else "")
        srcs = sorted(r[ab.SRCS])
        r[ab.IDX["sources"]] = ", ".join(srcs)
        r[ab.IDX["n_sources"]] = len(srcs)
    return rows


def vetted_keys(led):
    return {k for k, row in led.items() if any(row[c] for c in VET_COLS)}


def main():
    led = load_ledger()
    src = ab._argval("--harvest-from", "") or newest_queue()
    if "--harvest-from" in sys.argv and not os.path.exists(src):
        raise SystemExit(f"--harvest-from: {src} not found")
    if src and os.path.exists(src):
        harvest(led, src)
        save_ledger(led)
    else:
        print("no previous Vet_Queue workbook — first run, nothing to "
              "harvest", flush=True)
    if "--harvest-only" in sys.argv:
        print("--harvest-only: done", flush=True)
        return

    rows = load_rows()
    keep_keys = vetted_keys(led)
    if POSTED_SINCE:
        pre = len(rows)
        rows = [r for r in rows
                if not r[ab.IDX["posted"]]
                or r[ab.IDX["posted"]] >= POSTED_SINCE
                or r[ab.IDX["job_key"]] in keep_keys]
        print(f"posted-date cutoff {POSTED_SINCE}: cut {pre - len(rows)}, "
              f"kept {len(rows)}", flush=True)

    # ---- slice: US-remote + MA/RI/South-FL local
    queue = []            # (bucket, region, r)
    n_rem = n_loc = 0
    for r in rows:
        us = (r[ab.IDX["country"]] or "").strip().upper() in US_OK
        regs = ab.regions_of(r[ab.IDX["location"]], r[ab.IDX["state"]],
                             r[ab.IDX["country"]])
        if ab.is_remote(r) and us:
            queue.append(("Remote", "; ".join(regs), r))
            n_rem += 1
        elif r[ab.IDX["work_mode"]] in ("On-site", "", "Hybrid") and regs:
            queue.append(("Local", "; ".join(regs), r))
            n_loc += 1
    queue_keys = {r[ab.IDX["job_key"]] for _, _, r in queue}
    print(f"queue: {len(queue)} rows ({n_rem} US-remote, {n_loc} local)",
          flush=True)

    # ---- vetted-but-out-of-slice and vanished jobs stay visible.
    # Only the ledger's active, out-of-queue keys need a lookup; when the
    # ledger has none (e.g. nothing vetted yet), skip building the 1.5M-row
    # index entirely — it is the memory peak on a 16GB box.
    active_out = [(jk, lrow) for jk, lrow in led.items()
                  if jk not in queue_keys
                  and any(lrow[c] for c in VET_COLS)
                  and lrow["status"] not in INACTIVE]
    n_other = n_gone = 0
    if active_out:
        by_key = {r[ab.IDX["job_key"]]: r for r in rows}
        for jk, lrow in active_out:
            r = by_key.get(jk)
            if r is not None:
                regs = ab.regions_of(r[ab.IDX["location"]], r[ab.IDX["state"]],
                                     r[ab.IDX["country"]])
                queue.append(("Other", "; ".join(regs), r))
                n_other += 1
            else:
                ghost = [""] * ab.N
                for c in ("company", "title", "apply_url", "posted"):
                    ghost[ab.IDX[c]] = lrow[c]
                ghost[ab.IDX["job_key"]] = jk
                queue.append(("Gone", "", ghost))
                n_gone += 1
        del by_key
    if n_other or n_gone:
        print(f"kept visible: {n_other} active-vetted out-of-slice, "
              f"{n_gone} gone from crawl", flush=True)

    # queue tuples now hold refs to the ~78k rows we actually write; drop the
    # full 1.5M-row list so the other ~1.46M objects free before the workbook
    # write buffers up.
    del rows
    gc.collect()

    queue.sort(key=lambda t: t[2][ab.IDX["posted"]] or "", reverse=True)
    queue.sort(key=lambda t: (BUCKET_RANK[t[0]],
                              ab.REGION_RANK.get((t[1].split("; ") or [""])[0],
                                                 9)))

    seen = set()
    if os.path.exists(SEEN_KEYS):
        with open(SEEN_KEYS, encoding="utf-8") as f:
            seen = {ln.strip() for ln in f if ln.strip()}
    first_run = not seen

    # ---- summaries
    by_status, by_bucket, by_region = {}, {}, {}
    n_new = 0

    # ---- workbook (Queue FIRST: the dropdown splice targets sheet1.xml)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Queue")
    for j, wd in enumerate(QWIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = f"{get_column_letter(len(VET_COLS + CTX_COLS) + 1)}2"
    n_q = len(queue)
    last_col = get_column_letter(len(QCOLS))
    ws.auto_filter.ref = f"A1:{last_col}{n_q + 1}"
    for formula, color in STATUS_FILLS + [GONE_FILL]:
        ws.conditional_formatting.add(
            f"A2:{last_col}{n_q + 1}",
            FormulaRule(formula=[formula],
                        fill=PatternFill("solid", start_color=color,
                                         end_color=color),
                        stopIfTrue=False))
    ws.append(hdr(ws, QCOLS))
    for bucket, region, r in queue:
        jk = r[ab.IDX["job_key"]]
        lrow = led.get(jk)
        vet = [lrow[c] for c in VET_COLS] if lrow else [""] * len(VET_COLS)
        is_new = "" if first_run or bucket == "Gone" \
            else ("★" if jk not in seen else "")
        listing = "gone" if bucket == "Gone" else "live"
        ws.append(vet + [is_new, listing, bucket, region]
                  + [ab.safe(r[i]) for i in KEEP])
        st = (lrow["status"] if lrow else "") or "(unvetted)"
        by_status[st] = by_status.get(st, 0) + 1
        by_bucket[bucket] = by_bucket.get(bucket, 0) + 1
        if region:
            by_region[region] = by_region.get(region, 0) + 1
        if is_new:
            n_new += 1

    # ---- Top Roles: best AI-builder roles in the eligible slice, ranked
    # best-first (band desc, newest first within band), capped at TOP_N.
    scored = []
    for bucket, region, r in queue:
        if bucket == "Gone":
            continue
        band, kw = ai_band(r[ab.IDX["title"]] or "", r[ab.IDX["skills"]] or "")
        if band:
            scored.append((band, r[ab.IDX["posted"]] or "", kw, bucket,
                           region, r))
    scored.sort(key=lambda t: (t[0], t[1]), reverse=True)
    top = scored[:TOP_N]
    top_bands = {}
    for band, *_ in top:
        top_bands[AI_BANDS[band]] = top_bands.get(AI_BANDS[band], 0) + 1
    print(f"top roles: {len(top)} of {len(scored)} AI-builder rows in slice "
          f"({top_bands})", flush=True)

    # ---- Top Roles tab (sheet2 — dropdown splice targets it too)
    ws = wb.create_sheet("Top Roles")
    for j, wd in enumerate(TOPWIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = f"{get_column_letter(len(VET_COLS) + len(TOP_CTX) + 1)}2"
    n_top = len(top)
    top_last = get_column_letter(len(TOPCOLS))
    ws.auto_filter.ref = f"A1:{top_last}{n_top + 1}"
    for formula, color in STATUS_FILLS:      # no gone-fill: all top rows live
        ws.conditional_formatting.add(
            f"A2:{top_last}{n_top + 1}",
            FormulaRule(formula=[formula],
                        fill=PatternFill("solid", start_color=color,
                                         end_color=color),
                        stopIfTrue=False))
    ws.append(hdr(ws, TOPCOLS))
    for rank, (band, posted, kw, bucket, region, r) in enumerate(top, 1):
        jk = r[ab.IDX["job_key"]]
        lrow = led.get(jk)
        vet = [lrow[c] for c in VET_COLS] if lrow else [""] * len(VET_COLS)
        is_new = "" if first_run else ("★" if jk not in seen else "")
        ws.append(vet + [rank, AI_BANDS[band], kw, is_new, bucket, region]
                  + [ab.safe(r[i]) for i in KEEP])

    lines = [
        ("VET QUEUE — filter and vet Getro+Lever+Greenhouse jobs, "
         "decisions persist", ""),
        ("Snapshot (UTC)", SNAP),
        ("", ""),
        ("Queue rows", n_q),
        ("  US-remote", n_rem),
        ("  local (MA / RI / South-FL)", n_loc),
        ("  active-vetted, out of slice", n_other),
        ("  gone from crawl", n_gone),
        ("Top Roles (ranked AI-builder)", n_top),
        ("NEW since last queue", "n/a (first build)" if first_run else n_new),
        ("", ""),
        ("How to vet", "Pick a status in column A (dropdown: "
         + " / ".join(STATUSES) + "), jot notes/contact/dates in B–F. Rows "
         "color by status. Filter and sort freely — nothing you do to the "
         "view is lost, rows are matched by job_key, not position. The same "
         "job vetted on either the Queue or Top Roles tab carries to both."),
        ("Top Roles tab", f"The best ~{n_top} AI-building roles in this slice, "
         "ranked best-first: Claude / Anthropic / agentic / forward-deployed "
         "titles at the very top, then core AI engineer / LLM / GenAI, then "
         "AI-ML adjacent, then engineering titles whose skills list AI tooling "
         "— newest first within each band. 'ai_band' names the band, 'matched' "
         "the keyword that placed it, 'rank' its 1..N position. Rebuild with "
         "--top N to change how many."),
        ("Carry-forward", "Save this file, then rerun the build: every row "
         "with vet data is banked into vet_ledger.csv (by job_key) and "
         "re-applied to the fresh snapshot. Clearing all six vet cells on a "
         "row un-vets it. Vetted jobs never age out of the date cutoff."),
        ("new column", "★ = this job was not in any previous queue. Filter "
         "on it after a refresh to see only new arrivals."),
        ("listing / bucket", "'gone' = the posting left the crawl but your "
         "status is still active (Applied/Interviewing/...), so the row is "
         "kept from the ledger. Bucket 'Other' = a job you vetted that no "
         "longer matches the Remote/Local slice — still live, still shown."),
        ("What's in the queue", "US-remote jobs (remote + country US/blank) "
         "plus on-site/hybrid roles in MA, RI, or the Miami / Fort "
         "Lauderdale / West Palm Beach metros, from the full Getro + Lever "
         "+ Greenhouse universe (every live public Greenhouse board, added "
         "2026-07-21). Full references: Getro_Lever_Jobs_<date>.xlsx and "
         "Greenhouse_All_Jobs_<date>.xlsx."),
        ("Posted-date cutoff",
         (f"Jobs first seen before {POSTED_SINCE} are excluded (vetted rows "
          "exempt). --since YYYY-MM-DD moves the line, --all-dates disables.")
         if POSTED_SINCE else "None (--all-dates)."),
        ("Refresh", "python3 vet_build.py   (harvests this workbook, then "
         "rebuilds from the master CSVs). python3 vet_build.py "
         "--harvest-only just banks your vet work into the ledger."),
    ]
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 116
    bold = Font(bold=True)
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("Summary")
    for col, wd in zip("AB", (34, 14)):
        ws.column_dimensions[col].width = wd

    def section(title, d):
        ws.append(hdr(ws, [title, "jobs"]))
        for k, v in sorted(d.items(), key=lambda kv: -kv[1]):
            ws.append([k, v])
        ws.append([])

    section("by status", by_status)
    section("by bucket", by_bucket)
    section("by region", by_region)
    section("Top Roles by AI band", top_bands)
    ws.append([WriteOnlyCell(ws, value="NEW this snapshot"),
               "n/a" if first_run else n_new])
    ws.append([WriteOnlyCell(ws, value="TOP ROLES (ranked)"), n_top])
    ws.append([WriteOnlyCell(ws, value="TOTAL queue"), n_q])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)

    # ---- splice the status dropdown into both data sheets. In write-only
    # mode openpyxl assigns sheetN.xml in creation order, so Queue=sheet1,
    # Top Roles=sheet2.
    targets = {"xl/worksheets/sheet1.xml": n_q,
               "xl/worksheets/sheet2.xml": n_top}

    def dv(nrows):
        return (f'<dataValidations count="1"><dataValidation type="list" '
                f'allowBlank="1" showErrorMessage="0" sqref="A2:A{nrows + 1}">'
                f'<formula1>"{",".join(STATUSES)}"</formula1>'
                f'</dataValidation></dataValidations>')

    tmp = OUT_XLSX + ".tmp"
    with zipfile.ZipFile(OUT_XLSX) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in targets:
                xml = data.decode("utf-8")
                anchor = xml.rfind("</conditionalFormatting>")
                if anchor == -1:
                    raise SystemExit(f"{item.filename}: conditionalFormatting "
                                     "anchor missing")
                pos = anchor + len("</conditionalFormatting>")
                xml = xml[:pos] + dv(targets[item.filename]) + xml[pos:]
                data = xml.encode("utf-8")
                print(f"dropdown spliced into "
                      f"{os.path.basename(item.filename)}", flush=True)
            zout.writestr(item, data)
    os.replace(tmp, OUT_XLSX)
    with zipfile.ZipFile(OUT_XLSX) as z:
        for fn in targets:
            xml = z.read(fn).decode("utf-8")
            if xml.find("<dataValidations") < xml.rfind(
                    "</conditionalFormatting>"):
                raise SystemExit(f"{fn}: dropdown splice failed validation")
    print("dropdowns verified", flush=True)

    seen |= queue_keys
    tmp = SEEN_KEYS + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        f.write("\n".join(sorted(seen)) + "\n")
    os.replace(tmp, SEEN_KEYS)
    print(f"seen-keys: {len(seen)} total", flush=True)


if __name__ == "__main__":
    main()
