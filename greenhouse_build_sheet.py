#!/usr/bin/env python3
"""
greenhouse_build_sheet.py
Aggregate greenhouse_all_jobs.jsonl (raw per-board rows from
greenhouse_fetch_jobs.py) into a deduplicated master CSV + organized Excel
workbook (mirrors lever_build_sheet.py).

Dedupe, two passes:
  1. exact (region, job id) — alias tokens of the same board fetch the same
     postings twice (Greenhouse job ids are stable per posting)
  2. same (company, title, apply_url) — or (company, title, location) when
     no URL. Board tokens merged, earliest first_published kept.

Greenhouse specifics: no structured salary or commitment fields in the Job
Board API (pay lives in description prose only, which we don't store);
region is derived per job from its absolute_url (.eu.greenhouse.io -> eu).

Outputs (this directory):
  greenhouse_jobs_master.csv        full deduped table
  Greenhouse_All_Jobs_<date>.xlsx   README / All Jobs / Boards
"""

import csv
import json
import os
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

BASE = os.path.dirname(os.path.abspath(__file__))
IN_JSONL = os.path.join(BASE, "greenhouse_all_jobs.jsonl")
STATE_JSON = os.path.join(BASE, "greenhouse_fetch_state.json")
OUT_CSV = os.path.join(BASE, "greenhouse_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Greenhouse_All_Jobs_{SNAPSHOT}.xlsx")

XLSX_ROW_CAP = 1_048_000

COLUMNS = [
    "title", "company", "location", "work_mode", "department", "office",
    "posted", "updated", "requisition_id", "language", "apply_url",
    "board", "board_url", "region", "gh_job_id",
]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)
BOARDS, TS = N, N + 1  # aux tail slots

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


EU_HOST = ".eu.greenhouse.io"
REMOTE_RE = re.compile(r"remote|anywhere|work from home", re.I)
HYBRID_RE = re.compile(r"hybrid", re.I)


def ts_of(iso):
    if not iso:
        return 0
    try:
        return datetime.fromisoformat(iso).timestamp()
    except ValueError:
        try:
            return datetime.fromisoformat(
                iso.replace("Z", "+00:00")).timestamp()
        except ValueError:
            return 0


def norm(rec):
    j = rec["job"]
    token = rec["token"]
    loc = I(((j.get("location") or {}).get("name") or "").strip())
    offices = I("; ".join(o.get("name", "").strip()
                          for o in (j.get("offices") or [])
                          if o.get("name", "").strip()))
    depts = I("; ".join(d.get("name", "").strip()
                        for d in (j.get("departments") or [])
                        if d.get("name", "").strip()))
    url = j.get("absolute_url") or ""
    eu = EU_HOST in url
    if REMOTE_RE.search(loc) or REMOTE_RE.search(offices):
        wm = "Remote"
    elif HYBRID_RE.search(loc):
        wm = "Hybrid"
    else:
        wm = ""
    fp = j.get("first_published") or ""
    up = j.get("updated_at") or ""
    row = [""] * (N + 2)
    row[IDX["title"]] = safe((j.get("title") or "").strip())
    row[IDX["company"]] = I(safe(rec.get("company")
                                 or j.get("company_name") or token))
    row[IDX["location"]] = loc
    row[IDX["work_mode"]] = I(wm)
    row[IDX["department"]] = depts
    row[IDX["office"]] = offices
    row[IDX["posted"]] = I(fp[:10])
    row[IDX["updated"]] = I(up[:10])
    row[IDX["requisition_id"]] = safe(str(j.get("requisition_id") or ""))
    row[IDX["language"]] = I(j.get("language") or "")
    row[IDX["apply_url"]] = url
    row[IDX["board"]] = I(token)
    row[IDX["board_url"]] = I(
        f"https://job-boards{'.eu' if eu else ''}.greenhouse.io/{token}")
    row[IDX["region"]] = I("eu" if eu else "us")
    row[IDX["gh_job_id"]] = j.get("id") or ""
    row[BOARDS] = {I(token)}
    row[TS] = ts_of(fp)
    return row


MERGE_IF_EMPTY = [IDX[k] for k in ("apply_url", "location", "department",
                                   "office", "language", "requisition_id")]


def merge(a, b):
    a[BOARDS] |= b[BOARDS]
    if b[TS] and (not a[TS] or b[TS] < a[TS]):
        a[TS], a[IDX["posted"]] = b[TS], b[IDX["posted"]]
    for i in MERGE_IF_EMPTY:
        if a[i] == "" and b[i] != "":
            a[i] = b[i]


def main():
    by_id = {}
    raw_rows = 0
    bad = 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            raw_rows += 1
            if raw_rows % 100_000 == 0:
                print(f"  ...{raw_rows} raw rows", flush=True)
            try:
                r = norm(json.loads(line))
            except (json.JSONDecodeError, KeyError):
                bad += 1
                continue
            k = (r[IDX["region"]], r[IDX["gh_job_id"]])
            prev = by_id.get(k)
            if prev is not None:
                merge(prev, r)
            else:
                by_id[k] = r
    after_id = len(by_id)

    by_key = {}
    for r in by_id.values():
        k = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
             r[IDX["apply_url"]] or "loc:" + r[IDX["location"]].lower())
        prev = by_key.get(k)
        if prev is not None:
            merge(prev, r)
        else:
            by_key[k] = r
    by_id = None
    rows = sorted(by_key.values(), key=lambda r: r[TS], reverse=True)
    by_key = None
    print(f"raw rows {raw_rows} -> unique ids {after_id} -> deduped "
          f"{len(rows)}" + (f" ({bad} malformed skipped)" if bad else ""),
          flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print(f"wrote {OUT_CSV}", flush=True)

    # ---- board stats
    state = {}
    if os.path.exists(STATE_JSON):
        with open(STATE_JSON, encoding="utf-8") as f:
            state = json.load(f)
    board_stats = {}  # (board, region) -> [jobs, company, board_url, newest]
    for r in rows:
        k = (r[IDX["board"]], r[IDX["region"]])
        st = board_stats.setdefault(
            k, [0, r[IDX["company"]], r[IDX["board_url"]], ""])
        st[0] += 1
        if r[IDX["posted"]] > st[3]:
            st[3] = r[IDX["posted"]]

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def headers(ws, names):
        out = []
        for n_ in names:
            c = WriteOnlyCell(ws, value=n_)
            c.font = bold
            out.append(c)
        return out

    n_boards_fetched = sum(1 for v in state.values()
                           if v.get("status") == "done")
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Greenhouse job boards — full aggregation", ""),
        ("Snapshot (UTC)", SNAPSHOT),
        ("", ""),
        ("Boards fetched", n_boards_fetched),
        ("Boards with >=1 posting", len(board_stats)),
        ("Raw job rows", raw_rows),
        ("Unique jobs (deduped)", len(rows)),
        ("", ""),
        ("What this is", "Every posting on every live public Greenhouse job "
         "board we could discover (boards.greenhouse.io + "
         "job-boards.greenhouse.io + EU variants), found via the Wayback "
         "Machine URL index plus Greenhouse links inside the Getro dataset, "
         "then fetched from Greenhouse's public Job Board API."),
        ("How deduping works", "Same posting fetched twice (alias board "
         "tokens), or same company + title + apply URL, is one row."),
        ("Master CSV", "greenhouse_jobs_master.csv holds the complete "
         "table."),
        ("", ""),
        ("Column", "Meaning"),
        ("title / company / location", "Role, employer (board display "
         "name), posting location text"),
        ("work_mode", "Remote / Hybrid inferred from location & office "
         "text (Greenhouse has no structured field; blank = on-site or "
         "unspecified)"),
        ("department / office", "Greenhouse's org tags on the posting"),
        ("posted / updated", "first_published / updated_at dates (UTC)"),
        ("requisition_id", "Employer's internal req number, when set"),
        ("language", "Posting language code"),
        ("apply_url", "The hosted job page / application form"),
        ("board / board_url / region", "Greenhouse board token, its public "
         "board, us or eu datacenter"),
        ("gh_job_id", "Stable posting id inside Greenhouse"),
        ("Note on salary", "Greenhouse's public API has no structured pay "
         "fields — pay ranges only appear inside description prose, which "
         "this sheet doesn't store."),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    widths = [46, 26, 34, 10, 24, 24, 11, 11, 14, 8, 46, 22, 40, 7, 12]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    n_rows = min(len(rows), XLSX_ROW_CAP)
    ws.auto_filter.ref = f"A1:{get_column_letter(N)}{n_rows + 1}"
    ws.append(headers(ws, COLUMNS))
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])
    if len(rows) > XLSX_ROW_CAP:
        print(f"NOTE: {len(rows) - XLSX_ROW_CAP} rows beyond Excel cap "
              "left in CSV only", flush=True)

    ws = wb.create_sheet("Boards")
    for col, wdt in zip("ABCDEF", (28, 34, 7, 48, 10, 12)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(board_stats) + 1}"
    ws.append(headers(ws, ["board", "company", "region", "board_url",
                           "jobs", "newest_post"]))
    for (b, reg) in sorted(board_stats,
                           key=lambda k: -board_stats[k][0]):
        n_, comp, url, newest = board_stats[(b, reg)]
        ws.append([safe(b), safe(comp), reg, url, n_, newest])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
