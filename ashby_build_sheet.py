#!/usr/bin/env python3
"""
ashby_build_sheet.py
Aggregate ashby_all_jobs.jsonl (raw per-board rows from ashby_fetch_jobs.py)
into a deduplicated master CSV + organized Excel workbook (mirrors
greenhouse_build_sheet.py).

Dedupe, two passes:
  1. exact Ashby posting id — alias slugs of the same board fetch the same
     postings twice (ids are stable UUIDs per posting)
  2. same (company, title, apply_url) — or (company, title, location) when
     no URL. Board slugs merged, earliest posted kept.

Ashby specifics vs the other ATS pipes:
  - Structured compensation IS available (summaryComponents -> Salary
    min/max/currency/interval), so salary columns are real, not parsed prose.
  - employmentType is a machine value (FullTime/PartTime/Intern/Contract/
    Temporary) — mapped to our vocabulary.
  - isRemote + workplaceType (Remote/Hybrid/Onsite) give a reliable work_mode.
  - address.postalAddress yields city/region/country directly.
  - secondaryLocations holds extra locations (kept, joined, for the sheet).

Outputs (this directory):
  ashby_jobs_master.csv          full deduped table
  Ashby_All_Jobs_<date>.xlsx     README / All Jobs / Boards
"""

import csv
import json
import os
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
IN_JSONL = os.path.join(BASE, "ashby_all_jobs.jsonl")
BOARDS_CSV = os.path.join(BASE, "ashby_job_boards.csv")
OUT_CSV = os.path.join(BASE, "ashby_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Ashby_All_Jobs_{SNAPSHOT}.xlsx")

XLSX_ROW_CAP = 1_048_000

COLUMNS = [
    "title", "company", "location", "work_mode", "remote", "department",
    "team", "employment_type", "city", "state", "country",
    "secondary_locations", "posted", "salary_min", "salary_max",
    "salary_currency", "salary_period", "compensation_summary",
    "apply_url", "job_url", "board", "board_url", "ashby_job_id",
]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)
BOARDS, TS = N, N + 1  # aux tail slots

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


# Ashby employmentType -> our vocabulary
ET_MAP = {
    "fulltime": "Full-time", "parttime": "Part-time", "intern": "Internship",
    "contract": "Contract", "temporary": "Temp", "graduate": "Full-time",
}
REMOTE_RE = re.compile(r"remote|anywhere|work from home", re.I)
HYBRID_RE = re.compile(r"hybrid", re.I)
# 1 YEAR / 1 MONTH / 1 HOUR ... -> our period vocabulary
PERIOD_MAP = {"1 year": "year", "1 month": "month", "1 week": "week",
              "1 day": "day", "1 hour": "hour", "hourly": "hour"}


def ts_of(iso):
    if not iso:
        return 0
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).timestamp()
    except ValueError:
        return 0


def salary_of(comp):
    """-> (min, max, currency, period, summary) from Ashby compensation."""
    if not isinstance(comp, dict):
        return "", "", "", "", ""
    summary = comp.get("compensationTierSummary") or \
        comp.get("scrapeableCompensationSalarySummary") or ""
    for c in (comp.get("summaryComponents") or []):
        if (c.get("compensationType") or "").lower() != "salary":
            continue
        lo, hi = c.get("minValue"), c.get("maxValue")
        if lo is None and hi is None:
            continue
        interval = (c.get("interval") or "").strip().lower()
        return (str(lo) if lo is not None else "",
                str(hi) if hi is not None else "",
                c.get("currencyCode") or "",
                PERIOD_MAP.get(interval, interval.replace("1 ", "")),
                summary)
    return "", "", "", "", summary


def norm(rec):
    j = rec["job"]
    slug = rec["slug"]
    loc = I((j.get("location") or "").strip())
    wt = (j.get("workplaceType") or "").strip()
    if j.get("isRemote") or wt.lower() == "remote" or REMOTE_RE.search(loc):
        wm = "Remote"
    elif wt.lower() == "hybrid" or HYBRID_RE.search(loc):
        wm = "Hybrid"
    elif wt:
        wm = "On-site"
    else:
        wm = ""
    postal = ((j.get("address") or {}).get("postalAddress") or {})
    sec = "; ".join(
        (s.get("location") or "").strip()
        for s in (j.get("secondaryLocations") or [])
        if (s.get("location") or "").strip())
    lo, hi, cur, per, csum = salary_of(j.get("compensation"))
    pub = j.get("publishedAt") or ""
    et = ET_MAP.get((j.get("employmentType") or "").strip().lower(), "")

    row = [""] * (N + 2)
    row[IDX["title"]] = safe((j.get("title") or "").strip())
    row[IDX["company"]] = I(slug)
    row[IDX["location"]] = loc
    row[IDX["work_mode"]] = I(wm)
    row[IDX["remote"]] = "True" if wm == "Remote" else ""
    row[IDX["department"]] = I((j.get("department") or "").strip())
    row[IDX["team"]] = I((j.get("team") or "").strip())
    row[IDX["employment_type"]] = I(et)
    row[IDX["city"]] = I((postal.get("addressLocality") or "").strip())
    row[IDX["state"]] = I((postal.get("addressRegion") or "").strip())
    row[IDX["country"]] = I((postal.get("addressCountry") or "").strip())
    row[IDX["secondary_locations"]] = safe(sec)
    row[IDX["posted"]] = I(pub[:10])
    row[IDX["salary_min"]] = lo
    row[IDX["salary_max"]] = hi
    row[IDX["salary_currency"]] = I(cur)
    row[IDX["salary_period"]] = I(per)
    row[IDX["compensation_summary"]] = safe(csum)
    row[IDX["apply_url"]] = j.get("applyUrl") or ""
    row[IDX["job_url"]] = j.get("jobUrl") or ""
    row[IDX["board"]] = I(slug)
    row[IDX["board_url"]] = I(rec.get("board_url") or "")
    row[IDX["ashby_job_id"]] = j.get("id") or ""
    row[BOARDS] = {I(slug)}
    row[TS] = ts_of(pub)
    return row


def merge(a, b):
    """Fold b into a: keep the earliest posting, union the board list."""
    a[BOARDS] |= b[BOARDS]
    if b[TS] and (not a[TS] or b[TS] < a[TS]):
        a[IDX["posted"]] = b[IDX["posted"]]
        a[TS] = b[TS]
    for c in COLUMNS:  # fill blanks from the duplicate
        if not a[IDX[c]] and b[IDX[c]]:
            a[IDX[c]] = b[IDX[c]]


def main():
    csv.field_size_limit(10 ** 7)
    if not os.path.exists(IN_JSONL):
        raise SystemExit(f"missing {IN_JSONL} — run ashby_fetch_jobs.py")

    by_id, raw = {}, 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            raw += 1
            r = norm(rec)
            if not r[IDX["title"]]:
                continue
            jid = r[IDX["ashby_job_id"]]
            key = jid or f"{r[IDX['company']]}|{r[IDX['title']]}|" \
                         f"{r[IDX['location']]}"
            prev = by_id.get(key)
            if prev is None:
                by_id[key] = r
            else:
                merge(prev, r)
    n_ids = len(by_id)

    by_sig = {}
    for r in by_id.values():
        sig = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               (r[IDX["apply_url"]] or r[IDX["location"]]).lower())
        prev = by_sig.get(sig)
        if prev is None:
            by_sig[sig] = r
        else:
            merge(prev, r)
    rows = list(by_sig.values())
    by_id = by_sig = None

    for r in rows:
        names = sorted(r[BOARDS])
        r[IDX["board"]] = "; ".join(names[:6]) + (
            f" (+{len(names) - 6} more)" if len(names) > 6 else "")
    rows.sort(key=lambda r: (r[IDX["posted"]] or "", r[IDX["company"]]),
              reverse=True)

    print(f"raw rows {raw} -> unique ids {n_ids} -> deduped {len(rows)}",
          flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print(f"wrote {OUT_CSV}", flush=True)

    write_xlsx(rows)


def write_xlsx(rows):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    n_sal = sum(1 for r in rows if r[IDX["salary_min"]])
    n_rem = sum(1 for r in rows if r[IDX["remote"]])
    for a, b in [
        ("Ashby — all jobs", f"Snapshot {SNAPSHOT}"),
        ("Rows", f"{len(rows):,} deduplicated postings"),
        ("With salary", f"{n_sal:,} have structured pay from Ashby's "
                        f"compensation API (min/max/currency/period)"),
        ("Remote", f"{n_rem:,} flagged remote (isRemote / workplaceType)"),
        ("Source", "api.ashbyhq.com/posting-api/job-board/{slug}"
                   "?includeCompensation=true — one request per board, "
                   "all postings returned at once."),
        ("Boards", "ashby_job_boards.csv lists every live board and its job "
                   "count; ashby_candidates.csv is the full candidate set "
                   "before the live/dead split."),
        ("Company names", "Ashby's public API exposes no organization display "
                          "name, so company is the board slug. Downstream "
                          "(all_jobs_build.py) merges on the canonical "
                          "ats|ashby|org|id URL key, not on company."),
        ("Master CSV", "ashby_jobs_master.csv holds the complete table; this "
                       "workbook is capped at ~1.048M rows per sheet."),
    ]:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    ws.freeze_panes = "A2"
    widths = {"title": 52, "company": 24, "location": 30, "work_mode": 10,
              "remote": 8, "department": 22, "team": 20,
              "employment_type": 14, "city": 18, "state": 8, "country": 10,
              "secondary_locations": 34, "posted": 11, "salary_min": 11,
              "salary_max": 11, "salary_currency": 9, "salary_period": 9,
              "compensation_summary": 30, "apply_url": 46, "job_url": 46,
              "board": 22, "board_url": 40, "ashby_job_id": 38}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    hdr = []
    for c in COLUMNS:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for i, w_ in enumerate([28, 46, 56, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    if os.path.exists(BOARDS_CSV):
        with open(BOARDS_CSV, newline="", encoding="utf-8") as f:
            rd = csv.reader(f)
            head = next(rd, None)
            if head:
                cells = []
                for c in head:
                    cell = WriteOnlyCell(ws, value=c)
                    cell.font = bold
                    cells.append(cell)
                ws.append(cells)
            for row in rd:
                ws.append(row)

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
