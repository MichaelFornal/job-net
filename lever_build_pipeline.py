#!/usr/bin/env python3
"""
lever_build_pipeline.py
Build the personal pipeline workbook from lever_jobs_master.csv (same shape
as getro_build_pipeline.py):

  - every Remote role
  - every On-site / Hybrid / untagged role located in Massachusetts,
    Rhode Island, or the South Florida metros (Miami-Dade / Broward /
    Palm Beach), plus rows tagged only "Florida" statewide, labeled
    "FL (statewide tag)".

Lever locations are freeform text ("Boston", "Miami, FL", "Remote - US"),
so matching uses word-boundary regexes: unambiguous city names match bare
(miami, fort lauderdale, boca raton...), ambiguous ones require an FL tag
(hollywood, plantation, sunrise, jupiter...). Postings whose Lever country
code is set and isn't US never match a local region.

Adds flag + notes columns with an Interested / Maybe / No dropdown spliced
into the sheet XML (write-only openpyxl can't attach data validation);
"Interested" rows turn green. Output: Lever_Pipeline_<date>.xlsx
(Pipeline / Summary / README).
"""

import csv
import os
import re
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
IN_CSV = os.path.join(BASE, "lever_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Lever_Pipeline_{SNAPSHOT}.xlsx")

FLAG_CHOICES = "Interested,Maybe,No"

SFL_BARE = [
    "miami", "fort lauderdale", "boca raton", "coral gables",
    "delray beach", "boynton beach", "pompano beach", "deerfield beach",
    "hialeah", "doral", "aventura", "coral springs", "pembroke pines",
    "hallandale beach", "dania beach", "lauderhill", "tamarac",
    "opa-locka", "cutler bay", "palmetto bay", "sunny isles",
    "key biscayne", "wilton manors", "palm beach", "coconut creek",
    "cooper city", "juno beach", "tequesta", "miami-dade", "broward",
    "south florida",
]
SFL_NEED_FL = [
    "hollywood", "plantation", "sunrise", "davie", "weston", "miramar",
    "margate", "parkland", "homestead", "kendall", "jupiter", "wellington",
    "lake worth", "greenacres", "lantana", "riviera beach", "oakland park",
]

SFL_BARE_RE = re.compile(
    r"\b(" + "|".join(re.escape(c) for c in SFL_BARE) + r")\b")
SFL_FL_RE = re.compile(
    r"\b(" + "|".join(re.escape(c) for c in SFL_NEED_FL)
    + r"),?\s+(fl|florida)\b")
MA_RE = re.compile(r"massachusetts|,\s*ma\b|\bboston\b|greater boston")
RI_RE = re.compile(r"rhode island|,\s*ri\b|\bprovidence\b")

REGION_RANK = {"MA": 1, "RI": 2, "South FL": 3, "FL (statewide tag)": 4}


def statewide_fl(l):
    for seg in l.split(";"):
        first = seg.split(",", 1)[0].strip()
        if first in ("florida", "fl") or first.endswith("- florida") \
                or first.endswith("- fl"):
            return True
    return False


def regions_of(loc, country):
    if country and country != "US":
        return []
    l = loc.lower()
    out = []
    if MA_RE.search(l):
        out.append("MA")
    if RI_RE.search(l):
        out.append("RI")
    if SFL_BARE_RE.search(l) or SFL_FL_RE.search(l):
        out.append("South FL")
    elif "florida" in l and statewide_fl(l):
        out.append("FL (statewide tag)")
    return out


def main():
    csv.field_size_limit(10**7)
    with open(IN_CSV, newline="", encoding="utf-8") as f:
        r = csv.reader(f)
        master_cols = next(r)
        i_loc = master_cols.index("location")
        i_wm = master_cols.index("work_mode")
        i_posted = master_cols.index("posted")
        i_country = master_cols.index("country")

        rows = []
        n_remote = 0
        n_local = 0
        for row in r:
            wm = row[i_wm]
            regs = regions_of(row[i_loc], row[i_country])
            if wm == "Remote":
                rows.append((0, REGION_RANK.get(regs[0], 9) if regs else 9,
                             row[i_posted], "Remote", "; ".join(regs), row))
                n_remote += 1
            elif wm in ("On-site", "", "Hybrid") and regs:
                rows.append((1, REGION_RANK.get(regs[0], 9),
                             row[i_posted], "Local", "; ".join(regs), row))
                n_local += 1

    rows.sort(key=lambda t: t[2], reverse=True)
    rows.sort(key=lambda t: (t[0], t[1]))
    print(f"pipeline rows: {len(rows)} ({n_remote} remote, {n_local} local)",
          flush=True)

    cols = ["flag", "notes", "bucket", "region"] + master_cols
    ncols = len(cols)
    nrows = len(rows)

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("Pipeline")
    widths = [11, 24, 9, 18] + [46, 26, 34, 10, 12, 20, 20, 11, 11, 11, 9,
                                9, 8, 46, 46, 20, 36, 7, 38]
    for i, wdt in enumerate(widths[:ncols], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                        fill_type="solid")
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(ncols)}{nrows + 1}",
        FormulaRule(formula=['$A2="Interested"'], fill=green,
                    stopIfTrue=False))
    hdr = []
    for c in cols:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    for _, _, _, bucket, region, row in rows:
        ws.append(["", "", bucket, region] + row)

    ws = wb.create_sheet("Summary")
    counts = {}
    for _, _, _, bucket, region, row in rows:
        key = (bucket, region or "(anywhere)")
        counts[key] = counts.get(key, 0) + 1
    for col, wdt in zip("ABC", (12, 26, 12)):
        ws.column_dimensions[col].width = wdt
    ws.append([WriteOnlyCell(ws, value=v) for v in ("bucket", "region",
                                                    "jobs")])
    for (bucket, region), n in sorted(counts.items(), key=lambda kv: -kv[1]):
        ws.append([bucket, region, n])
    ws.append(["TOTAL", "", nrows])

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Personal pipeline — Lever", f"Snapshot {SNAPSHOT} (UTC). {nrows} "
         "roles: every Remote job across all live public Lever boards, "
         "plus every on-site/hybrid job in MA, RI, and the Miami / Fort "
         "Lauderdale / West Palm Beach metros."),
        ("flag", f"Pick {FLAG_CHOICES.replace(',', ' / ')} from the "
         "dropdown (free text works too). Rows flagged Interested turn "
         "green. Use the filter arrow on 'flag' to see only your picks."),
        ("notes", "Free text — next steps, contacts, application status."),
        ("bucket", "Remote (works from anywhere) or Local (on-site in "
         "your regions)."),
        ("region", "MA / RI / South FL. 'FL (statewide tag)' = the "
         "posting says only 'Florida' with no city — could be your metro, "
         "could be Jacksonville; filter it out via the region column if "
         "it adds noise. Remote rows show a region when the employer also "
         "tagged a local office."),
        ("caveat", "Lever locations are freeform text, so region matching "
         "is best-effort: unambiguous cities (Miami, Boca Raton, Boston) "
         "match bare; ambiguous ones (Hollywood, Sunrise, Jupiter) only "
         "match with an FL tag. A posting whose country code isn't US "
         "never matches a local region."),
        ("everything else", "Columns carried from lever_jobs_master.csv — "
         "see the README in Lever_All_Jobs for the full glossary."),
        ("refresh", "Re-run lever_fetch_jobs.py + lever_build_sheet.py, "
         "then lever_build_pipeline.py. Re-apply your flags to a new "
         "snapshot by matching lever_job_id."),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)

    # ---- splice the flag dropdown into the Pipeline sheet XML
    dv_xml = (f'<dataValidations count="1">'
              f'<dataValidation type="list" allowBlank="1" '
              f'showErrorMessage="0" sqref="A2:A{nrows + 1}">'
              f'<formula1>"{FLAG_CHOICES}"</formula1>'
              f'</dataValidation></dataValidations>')
    tmp = OUT_XLSX + ".tmp"
    with zipfile.ZipFile(OUT_XLSX) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                anchor = xml.rfind("</conditionalFormatting>")
                if anchor == -1:
                    raise SystemExit("conditionalFormatting anchor missing")
                pos = anchor + len("</conditionalFormatting>")
                xml = xml[:pos] + dv_xml + xml[pos:]
                data = xml.encode("utf-8")
                print("dropdown spliced into sheet1", flush=True)
            zout.writestr(item, data)
    os.replace(tmp, OUT_XLSX)

    with zipfile.ZipFile(OUT_XLSX) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    i_dv = xml.find("<dataValidations")
    i_cf = xml.rfind("</conditionalFormatting>")
    i_pm = xml.find("<pageMargins")
    ok = (i_dv != -1 and i_cf != -1 and i_dv > i_cf
          and (i_pm == -1 or i_dv < i_pm)
          and xml.count("<dataValidations") == 1)
    print(f"verify: dropdown present and schema-ordered: {ok}", flush=True)
    if not ok:
        raise SystemExit("dataValidations splice verification failed")


if __name__ == "__main__":
    main()
