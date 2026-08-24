#!/usr/bin/env python3
"""
lever_build_sheet.py
Aggregate lever_all_jobs.jsonl (raw per-board rows from lever_fetch_jobs.py)
into a deduplicated master CSV + organized Excel workbook.

Dedupe, two passes (mirrors getro_build_sheet.py):
  1. exact (region, posting id) — case-variant slugs of the same board fetch
     the same postings twice
  2. same (company, title, apply_url) — or (company, title, locations) when
     no URL. Board slugs merged, earliest createdAt kept.

Outputs (this directory):
  lever_jobs_master.csv        full deduped table
  Lever_All_Jobs_<date>.xlsx   README / All Jobs / Boards
"""

import csv
import json
import os
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

BASE = os.path.dirname(os.path.abspath(__file__))
IN_JSONL = os.path.join(BASE, "lever_all_jobs.jsonl")
STATE_JSON = os.path.join(BASE, "lever_fetch_state.json")
OUT_CSV = os.path.join(BASE, "lever_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, f"Lever_All_Jobs_{SNAPSHOT}.xlsx")

XLSX_ROW_CAP = 1_048_000

COLUMNS = [
    "title", "company", "location", "work_mode", "commitment", "department",
    "team", "posted", "salary_min", "salary_max", "salary_currency",
    "salary_period", "country", "apply_url", "posting_url", "board",
    "board_url", "region", "lever_job_id",
]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)
BOARDS, TS = N, N + 1  # aux tail slots

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


WM = {"remote": "Remote", "hybrid": "Hybrid", "on-site": "On-site",
      "onsite": "On-site", "unspecified": "", "": ""}

# probe-era junk titles (hidden boards serve a 404 page while the API is
# live); fall back to the slug
BAD_TITLE_RE = re.compile(
    r"not found|404|\berror\b|just a moment|attention required|"
    r"access denied|page unavailable", re.I)

PERIOD = {"per-year-salary": "Year", "per-hour-wage": "Hour",
          "per-month-salary": "Month", "per-week-salary": "Week",
          "per-day-wage": "Day"}


def pretty_period(iv):
    if not iv:
        return ""
    if iv in PERIOD:
        return PERIOD[iv]
    return I(iv.replace("per-", "").replace("-salary", "").replace(
        "-wage", "").replace("-", " ").strip().title())


def norm(rec):
    j = rec["job"]
    cats = j.get("categories") or {}
    locs = cats.get("allLocations") or (
        [cats["location"]] if cats.get("location") else [])
    loc = I("; ".join(x.strip() for x in locs if x and x.strip()))
    wm = WM.get((j.get("workplaceType") or "").lower(),
                (j.get("workplaceType") or "").title())
    if not wm and "remote" in loc.lower():
        wm = "Remote"
    ts = j.get("createdAt")
    ts = ts / 1000 if isinstance(ts, (int, float)) and ts > 0 else 0
    sal = j.get("salaryRange") or {}
    mn, mx = sal.get("min"), sal.get("max")
    row = [""] * (N + 2)
    row[IDX["title"]] = safe((j.get("text") or "").strip())
    comp = rec.get("company") or rec["slug"]
    if BAD_TITLE_RE.search(comp):
        comp = rec["slug"]
    row[IDX["company"]] = I(comp)
    row[IDX["location"]] = loc
    row[IDX["work_mode"]] = I(wm)
    row[IDX["commitment"]] = I(cats.get("commitment") or "")
    row[IDX["department"]] = I(cats.get("department") or "")
    row[IDX["team"]] = I(cats.get("team") or "")
    row[IDX["posted"]] = I(datetime.fromtimestamp(
        ts, tz=timezone.utc).strftime("%Y-%m-%d")) if ts else ""
    row[IDX["salary_min"]] = mn if isinstance(mn, (int, float)) else ""
    row[IDX["salary_max"]] = mx if isinstance(mx, (int, float)) else ""
    row[IDX["salary_currency"]] = I((sal.get("currency") or "").upper())
    row[IDX["salary_period"]] = pretty_period(sal.get("interval") or "")
    row[IDX["country"]] = I(j.get("country") or "")
    row[IDX["apply_url"]] = j.get("applyUrl") or ""
    row[IDX["posting_url"]] = j.get("hostedUrl") or ""
    row[IDX["board"]] = I(rec["slug"])
    row[IDX["board_url"]] = I(rec.get("board_url") or "")
    row[IDX["region"]] = I(rec.get("region") or "us")
    row[IDX["lever_job_id"]] = j.get("id") or ""
    row[BOARDS] = {I(rec["slug"])}
    row[TS] = ts
    return row


MERGE_IF_EMPTY = [IDX[k] for k in ("apply_url", "location", "department",
                                   "team", "salary_min", "salary_max",
                                   "country")]


def merge(a, b):
    a[BOARDS] |= b[BOARDS]
    if b[TS] and (not a[TS] or b[TS] < a[TS]):
        a[TS], a[IDX["posted"]] = b[TS], b[IDX["posted"]]
    for i in MERGE_IF_EMPTY:
        if a[i] == "" and b[i] != "":
            a[i] = b[i]


def main():
    by_id = {}
    raw_rows = 0
    bad = 0
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            raw_rows += 1
            if raw_rows % 100_000 == 0:
                print(f"  ...{raw_rows} raw rows", flush=True)
            try:
                r = norm(json.loads(line))
            except (json.JSONDecodeError, KeyError):
                bad += 1
                continue
            k = (r[IDX["region"]], r[IDX["lever_job_id"]])
            prev = by_id.get(k)
            if prev is not None:
                merge(prev, r)
            else:
                by_id[k] = r
    after_id = len(by_id)

    by_key = {}
    for r in by_id.values():
        k = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
             r[IDX["apply_url"]] or "loc:" + r[IDX["location"]].lower())
        prev = by_key.get(k)
        if prev is not None:
            merge(prev, r)
        else:
            by_key[k] = r
    by_id = None
    rows = sorted(by_key.values(), key=lambda r: r[TS], reverse=True)
    by_key = None
    print(f"raw rows {raw_rows} -> unique ids {after_id} -> deduped "
          f"{len(rows)}" + (f" ({bad} malformed skipped)" if bad else ""),
          flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print(f"wrote {OUT_CSV}", flush=True)

    # ---- board stats
    state = {}
    if os.path.exists(STATE_JSON):
        with open(STATE_JSON, encoding="utf-8") as f:
            state = json.load(f)
    board_stats = {}  # (board, region) -> [jobs, company, board_url, newest]
    for r in rows:
        k = (r[IDX["board"]], r[IDX["region"]])
        st = board_stats.setdefault(
            k, [0, r[IDX["company"]], r[IDX["board_url"]], ""])
        st[0] += 1
        if r[IDX["posted"]] > st[3]:
            st[3] = r[IDX["posted"]]

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def headers(ws, names):
        out = []
        for n_ in names:
            c = WriteOnlyCell(ws, value=n_)
            c.font = bold
            out.append(c)
        return out

    n_boards_fetched = sum(1 for v in state.values()
                           if v.get("status") == "done")
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Lever job boards — full aggregation", ""),
        ("Snapshot (UTC)", SNAPSHOT),
        ("", ""),
        ("Boards fetched", n_boards_fetched),
        ("Boards with >=1 posting", len(board_stats)),
        ("Raw job rows", raw_rows),
        ("Unique jobs (deduped)", len(rows)),
        ("", ""),
        ("What this is", "Every posting on every live public Lever job "
         "board we could discover (jobs.lever.co + jobs.eu.lever.co), "
         "found via the Wayback Machine URL index plus Lever apply links "
         "inside the Getro dataset, then fetched from Lever's public "
         "postings API."),
        ("How deduping works", "Same posting fetched twice (case-variant "
         "board slugs), or same company + title + apply URL, is one row."),
        ("Master CSV", "lever_jobs_master.csv holds the complete table."),
        ("", ""),
        ("Column", "Meaning"),
        ("title / company / location", "Role, employer (board display "
         "name), posting location(s)"),
        ("work_mode", "On-site / Hybrid / Remote (Lever's workplaceType; "
         "blank = unspecified)"),
        ("commitment", "Full-time / Part-time / Contract / Intern etc."),
        ("department / team", "Lever's org tags on the posting"),
        ("posted", "Posting creation date (UTC)"),
        ("salary_min / salary_max", "Published salary range in "
         "salary_currency per salary_period; blank when the employer "
         "doesn't publish structured pay"),
        ("country", "ISO country code when Lever provides one"),
        ("apply_url / posting_url", "Direct application form / the hosted "
         "job page"),
        ("board / board_url / region", "Lever account slug, its public "
         "board, us or eu datacenter"),
        ("lever_job_id", "Stable posting UUID inside Lever"),
    ]
    for a, b in lines:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    widths = [46, 26, 34, 10, 12, 20, 20, 11, 11, 11, 9, 9, 8, 46, 46, 20,
              36, 7, 38]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    n_rows = min(len(rows), XLSX_ROW_CAP)
    ws.auto_filter.ref = f"A1:{get_column_letter(N)}{n_rows + 1}"
    ws.append(headers(ws, COLUMNS))
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])
    if len(rows) > XLSX_ROW_CAP:
        print(f"NOTE: {len(rows) - XLSX_ROW_CAP} rows beyond Excel cap "
              "left in CSV only", flush=True)

    ws = wb.create_sheet("Boards")
    for col, wdt in zip("ABCDEF", (28, 34, 7, 44, 10, 12)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(board_stats) + 1}"
    ws.append(headers(ws, ["board", "company", "region", "board_url",
                           "jobs", "newest_post"]))
    for (b, reg) in sorted(board_stats,
                           key=lambda k: -board_stats[k][0]):
        n_, comp, url, newest = board_stats[(b, reg)]
        ws.append([safe(b), safe(comp), reg, url, n_, newest])

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}", flush=True)


if __name__ == "__main__":
    main()
