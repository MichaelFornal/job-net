#!/usr/bin/env python3
"""
yc_build_sheet.py — yc_all_jobs.jsonl -> yc_jobs_master.csv +
YC_All_Jobs_<date>.xlsx.

Two column choices are load-bearing:

  apply_url = job_url = https://www.ycombinator.com/companies/{slug}/jobs/
                        {hashid}-{title-slug}
    This is the URL all_jobs_build.canon_key parses into
    ats|ycwaas|{slug}|{hashid}.  job_key() tries canon_key(apply_url) FIRST,
    so putting the hashid URL here is what merges these rows with the copies
    Getro already holds.

  signup_url = the account.ycombinator.com/authenticate?...signup_job_id=N
    link, kept for the human but deliberately NOT mapped to apply_url.  It
    carries a different id (numeric) from the hashid in the page URL, with no
    public mapping between them, so treating it as the canonical URL would
    give one posting two disagreeing tier-1 keys — worse than none, because
    tiers never cross-match.

posted comes from yc_posted_ledger.json, never recomputed here: `createdAt` is
a relative string and re-deriving it each run would walk every row backwards
until it fell off the cutoff.
"""
import csv
import json
import os
import re
import sys
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

BASE = os.path.dirname(os.path.abspath(__file__))
SMOKE = int(os.environ.get("SMOKE") or 0)
PFX = "smoke_" if SMOKE else ""

IN_JSONL = os.path.join(BASE, PFX + "yc_all_jobs.jsonl")
LEDGER_JSON = os.path.join(BASE, PFX + "yc_posted_ledger.json")
COMPANIES_JSON = os.path.join(BASE, PFX + "yc_companies.json")
OUT_CSV = os.path.join(BASE, PFX + "yc_jobs_master.csv")
SNAPSHOT = datetime.now(timezone.utc).strftime("%Y-%m-%d")
OUT_XLSX = os.path.join(BASE, "%sYC_All_Jobs_%s.xlsx" % (
    "Smoke_" if SMOKE else "", SNAPSHOT))
XLSX_ROW_CAP = 1_048_000

COLUMNS = ["title", "company", "location", "work_mode", "remote",
           "employment_type", "department", "role_type", "min_experience",
           "visa", "posted", "posted_precision", "created_raw",
           "salary_min", "salary_max", "salary_currency", "salary_period",
           "salary_range_raw", "equity", "company_batch", "company_stage",
           "company_size", "company_industry", "company_website",
           "apply_url", "signup_url", "job_url", "board", "board_url",
           "yc_job_id"]
IDX = {c: i for i, c in enumerate(COLUMNS)}
N = len(COLUMNS)

REMOTE_RE = re.compile(r"remote|anywhere|work from home", re.I)
HYBRID_RE = re.compile(r"hybrid", re.I)
ET_MAP = {"intern": "Internship", "internship": "Internship"}
# "$90K - $120K", "€70K - €90K", "$150,000 - $200,000"
MONEY_RE = re.compile(r"([$€£]?)\s*([\d][\d,.]*)\s*([KkMm])?")
CUR = {"$": "USD", "€": "EUR", "£": "GBP"}

_MEMO = {}


def I(s):
    return _MEMO.setdefault(s, s)


def safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def parse_salary(text):
    """-> (min, max, currency).  Blank rather than wrong: a pair that parses
    below 1000 with no K/M suffix is an hourly rate or a typo, and writing it
    into a column ai_ma_remote.py sorts on would be worse than leaving it
    empty (salary_range_raw keeps the original either way)."""
    if not text:
        return "", "", ""
    vals, cur = [], ""
    for m in MONEY_RE.finditer(str(text)):
        sym, num, suf = m.group(1), m.group(2), m.group(3)
        try:
            v = float(num.replace(",", ""))
        except ValueError:
            continue
        if suf and suf.lower() == "k":
            v *= 1_000
        elif suf and suf.lower() == "m":
            v *= 1_000_000
        if sym and not cur:
            cur = CUR.get(sym, "")
        vals.append(v)
    if not vals:
        return "", "", ""
    lo, hi = min(vals), max(vals)
    if hi < 1000:
        return "", "", ""
    return str(int(lo)), str(int(hi)), cur or "USD"


def norm(rec, ledger, companies):
    j = rec.get("job") or {}
    slug = rec.get("slug") or ""
    yc = rec.get("yc") or {}
    meta = companies.get(slug) or {}
    jid = str(j.get("id") or "").strip()
    title = (j.get("title") or "").strip()
    loc = I((j.get("location") or "").strip())

    # No structured remote field exists; infer from the free-text location.
    # Deliberately NOT from the company-level yc-oss "regions", which says
    # something about the company, not this job.
    if REMOTE_RE.search(loc):
        wm = "Remote"
    elif HYBRID_RE.search(loc):
        wm = "Hybrid"
    elif loc:
        wm = "On-site"
    else:
        wm = ""

    et_raw = (j.get("type") or "").strip()
    et = ET_MAP.get(et_raw.lower(), et_raw)

    path = (j.get("url") or "").strip()
    job_url = ("https://www.ycombinator.com" + path) if path.startswith("/") \
        else path
    lo, hi, cur = parse_salary(j.get("salaryRange"))
    entry = ledger.get(jid) or {}

    r = [""] * N
    r[IDX["title"]] = title
    r[IDX["company"]] = I(j.get("companyName") or yc.get("name")
                          or meta.get("company") or slug)
    r[IDX["location"]] = loc
    r[IDX["work_mode"]] = I(wm)
    r[IDX["remote"]] = "True" if wm == "Remote" else ""
    r[IDX["employment_type"]] = I(et)
    r[IDX["department"]] = I((j.get("prettyRole") or j.get("role") or "").strip())
    r[IDX["role_type"]] = I((j.get("roleSpecificType") or "") or "")
    r[IDX["min_experience"]] = I((j.get("minExperience") or "") or "")
    r[IDX["visa"]] = I((j.get("visa") or "") or "")
    r[IDX["posted"]] = I(entry.get("posted") or "")
    r[IDX["posted_precision"]] = I(entry.get("precision") or "")
    r[IDX["created_raw"]] = I(entry.get("created_raw")
                              or (j.get("createdAt") or ""))
    r[IDX["salary_min"]] = lo
    r[IDX["salary_max"]] = hi
    r[IDX["salary_currency"]] = I(cur)
    r[IDX["salary_period"]] = I("year" if lo else "")
    r[IDX["salary_range_raw"]] = I((j.get("salaryRange") or "") or "")
    r[IDX["equity"]] = I((j.get("equityRange") or "") or "")
    r[IDX["company_batch"]] = I(j.get("companyBatchName")
                                or yc.get("batch_name") or meta.get("batch")
                                or "")
    r[IDX["company_stage"]] = I(meta.get("stage") or "")
    r[IDX["company_size"]] = I(str(yc.get("team_size")
                                   or meta.get("team_size") or ""))
    r[IDX["company_industry"]] = I(meta.get("industry") or "")
    r[IDX["company_website"]] = I(yc.get("website") or meta.get("website")
                                  or "")
    r[IDX["apply_url"]] = job_url
    r[IDX["signup_url"]] = (j.get("applyUrl") or "")
    r[IDX["job_url"]] = job_url
    r[IDX["board"]] = I(slug)
    r[IDX["board_url"]] = I(rec.get("board_url") or "")
    r[IDX["yc_job_id"]] = jid
    return r


def main():
    if not os.path.exists(IN_JSONL):
        raise SystemExit("missing %s — run yc_fetch_jobs.py" % IN_JSONL)
    ledger, companies = {}, {}
    for path, tgt in ((LEDGER_JSON, "ledger"), (COMPANIES_JSON, "companies")):
        if os.path.exists(path):
            try:
                with open(path, encoding="utf-8") as f:
                    d = json.load(f)
            except (OSError, json.JSONDecodeError):
                d = {}
            if tgt == "ledger":
                ledger = d
            else:
                companies = d

    by_id, raw = {}, 0
    hash_of = {}
    with open(IN_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            raw += 1
            r = norm(rec, ledger, companies)
            if not r[IDX["title"]]:
                continue
            jid = r[IDX["yc_job_id"]]
            key = jid or "%s|%s" % (r[IDX["company"]], r[IDX["title"]])
            by_id.setdefault(key, r)
            # hashid must be injective w.r.t. the numeric id, or the
            # canon_key branch would merge two different postings.
            u = r[IDX["job_url"]].rsplit("/", 1)[-1].split("-")[0].lower()
            if u:
                prev = hash_of.setdefault((r[IDX["board"]], u), jid)
                if prev != jid:
                    print("  WARNING: hashid collision %s/%s -> ids %s and %s"
                          % (r[IDX["board"]], u, prev, jid), flush=True)
    n_ids = len(by_id)

    by_sig = {}
    for r in by_id.values():
        sig = (r[IDX["company"]].lower(), r[IDX["title"]].lower(),
               (r[IDX["apply_url"]] or r[IDX["location"]]).lower())
        by_sig.setdefault(sig, r)
    rows = list(by_sig.values())
    rows.sort(key=lambda r: (r[IDX["posted"]] or "", r[IDX["company"]]),
              reverse=True)
    print("raw rows %d -> unique ids %d -> deduped %d"
          % (raw, n_ids, len(rows)), flush=True)

    n_fs = sum(1 for r in rows if r[IDX["posted_precision"]] == "first_seen")
    if rows and n_fs / len(rows) > 0.05:
        print("  WARNING: %d rows (%.1f%%) fell back to first_seen dates — "
              "the createdAt vocabulary may have changed"
              % (n_fs, 100.0 * n_fs / len(rows)), flush=True)

    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow(r[:N])
    print("wrote %s" % OUT_CSV, flush=True)
    write_xlsx(rows)


def write_xlsx(rows):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    n_sal = sum(1 for r in rows if r[IDX["salary_min"]])
    n_eq = sum(1 for r in rows if r[IDX["equity"]])
    for a, b in [
        ("Y Combinator — Work at a Startup", "Snapshot %s" % SNAPSHOT),
        ("Rows", "%s deduplicated postings" % format(len(rows), ",")),
        ("With salary", "%s carry a salary range; %s carry equity"
         % (format(n_sal, ","), format(n_eq, ","))),
        ("Source", "www.ycombinator.com/companies/{slug}/jobs with "
                   "Accept: text/html — the job list is inline in the "
                   "Inertia data-page attribute. Accept: application/json "
                   "returns 404 on the same URL."),
        ("Slugs", "yc-oss.github.io/api/companies/hiring.json (refreshed "
                  "daily), plus a grep of the local corpus."),
        ("posted", "Derived from YC's relative createdAt string ('5 days', "
                   "'about 2 months') and FROZEN in yc_posted_ledger.json at "
                   "first sight. Recomputing each run would age every row "
                   "backwards until it fell off the cutoff while still live."),
        ("apply_url", "The hashid page URL, which is what canon_key merges "
                      "on. signup_url holds the account.ycombinator.com "
                      "link, which carries a different id namespace and is "
                      "kept out of the dedupe path on purpose."),
        ("Overlap", "WaaS postings are not syndicated to any ATS, so most of "
                    "this is net-new against the rest of the corpus."),
    ]:
        ca = WriteOnlyCell(ws, value=a)
        ca.font = bold
        cb = WriteOnlyCell(ws, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([ca, cb])

    ws = wb.create_sheet("All Jobs")
    ws.freeze_panes = "A2"
    widths = {"title": 50, "company": 24, "location": 28, "work_mode": 10,
              "remote": 8, "employment_type": 14, "department": 18,
              "role_type": 16, "min_experience": 20, "visa": 22,
              "posted": 11, "posted_precision": 15, "created_raw": 14,
              "salary_min": 11, "salary_max": 11, "salary_currency": 9,
              "salary_period": 9, "salary_range_raw": 18, "equity": 14,
              "company_batch": 12, "company_stage": 12, "company_size": 11,
              "company_industry": 18, "company_website": 30,
              "apply_url": 56, "signup_url": 46, "job_url": 56,
              "board": 22, "board_url": 46, "yc_job_id": 10}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    hdr = []
    for c in COLUMNS:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(N),
                                      min(len(rows), XLSX_ROW_CAP) + 1)
    for r in rows[:XLSX_ROW_CAP]:
        ws.append([safe(v) for v in r[:N]])

    ws = wb.create_sheet("Boards")
    for col, wd in zip("ABCDE", (26, 26, 10, 12, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    hdr = []
    for c in ("board", "company", "jobs", "batch", "newest_post"):
        cell = WriteOnlyCell(ws, value=c)
        cell.font = bold
        hdr.append(cell)
    ws.append(hdr)
    stats = {}
    for r in rows:
        st = stats.setdefault(r[IDX["board"]],
                              [r[IDX["company"]], 0, r[IDX["company_batch"]],
                               ""])
        st[1] += 1
        if (r[IDX["posted"]] or "") > st[3]:
            st[3] = r[IDX["posted"]] or ""
    for b, (comp, n, batch, newest) in sorted(stats.items(),
                                              key=lambda kv: -kv[1][1]):
        ws.append([safe(b), safe(comp), n, safe(batch), newest])

    wb.save(OUT_XLSX)
    print("wrote %s" % OUT_XLSX, flush=True)


if __name__ == "__main__":
    sys.exit(main())
